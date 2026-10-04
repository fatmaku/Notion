# Für die Gegenprüfung: neue Fassung mit IDs + Hunk-Liste (alt → neu) mit Kontext.
import difflib, sys
from docx import Document
key, old, new, outdir = sys.argv[1:5]
A = [p.text for p in Document(old).paragraphs]
P = Document(new).paragraphs
B = [p.text for p in P]
with open(f"{outdir}/{key}_neu.txt", "w") as f:
    for i, p in enumerate(P):
        st = p.style.name
        tag = f"[{i}]" + (f" <{st}>" if st != "Normal" else "")
        f.write(f"{tag} {p.text if p.text.strip() else '(boş)'}\n")
sm = difflib.SequenceMatcher(None, A, B, autojunk=False)
out = []; n = 0
for tag, i1, i2, j1, j2 in sm.get_opcodes():
    if tag == "equal": continue
    n += 1
    ctx_before = next((f"[{j}] {B[j][:300]}" for j in range(j1 - 1, -1, -1) if B[j].strip()), "")
    ctx_after = next((f"[{j}] {B[j][:300]}" for j in range(j2, len(B)) if B[j].strip()), "")
    out.append(f"=== Hunk {n}: {tag} (alt {i1}-{i2-1} → neu {j1}-{j2-1})")
    out.append(f"  KONTEXT DAVOR: {ctx_before}")
    for k in range(i1, i2):
        if A[k].strip(): out.append(f"  - alt[{k}] {A[k]}")
    for k in range(j1, j2):
        if B[k].strip(): out.append(f"  + neu[{k}] {B[k]}")
    out.append(f"  KONTEXT DANACH: {ctx_after}")
open(f"{outdir}/{key}_hunks.txt", "w").write("\n".join(out) + "\n")
print(key, n, "Hunks;", len(B), "Absätze neu")
