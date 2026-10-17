#!/usr/bin/env python3
"""BEN YOKSAM – zusammengefügte Units → DOCX (Absatzvorlagen und Bilder aus YOLCU_TR_v5.docx als Vorlage).

  python3 tools/byk_build.py [--out src/BEN_YOKSAM_TR_v1.docx]
Liest benyoksam/BEN_YOKSAM_units.json (von byk_check.py --all geschrieben).
"""
import argparse, copy, json, sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from i18n_common import parse_markup, strip_tags  # noqa: E402

I = Path(__file__).resolve().parent.parent
TEMPLATE = I / "src" / "YOLCU_TR_v5.docx"
IMAGE_PIDS = [3, 35, 1238, 3384]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(I / "src" / "BEN_YOKSAM_TR_v1.docx"))
    ap.add_argument("--units", default=str(I / "benyoksam" / "BEN_YOKSAM_units.json"))
    a = ap.parse_args()
    paras = json.load(open(a.units, encoding="utf-8"))["paras"]
    doc = Document(TEMPLATE)
    src_paras = doc.paragraphs
    img_xml = {pid: copy.deepcopy(src_paras[pid]._element) for pid in IMAGE_PIDS}
    # rPr-Vorlage je Stil (Schrift usw. wie im Original), ohne Fett/Kursiv
    rpr = {}
    for p in src_paras:
        st = p.style.name
        if st in rpr:
            continue
        for r in p.runs:
            el = r._element.find(qn("w:rPr"))
            if el is not None:
                t = copy.deepcopy(el)
                for tag in ("w:b", "w:bCs", "w:i", "w:iCs"):
                    for x in t.findall(qn(tag)):
                        t.remove(x)
                rpr[st] = t
                break
    # Absatzvorlagen der Zwischenkapitel (Tuncays „ÜÇÜNCÜ GÖZ“) anlegen, falls die Vorlage sie nicht hat
    from docx.enum.style import WD_STYLE_TYPE
    names = {st.name for st in doc.styles}
    # Alle auf „Normal“ (keine Überschriften-Vererbung): „Interlude Title“ bekommt Gliederungsebene 1 (wie die Kapitel, nicht deren
    # Unterpunkt im Navigationsbereich/Inhaltsverzeichnis), „Interlude Heading“ bleibt Textkörper. Schrift (rPr) wie die Vorlage „look“.
    from docx.oxml import OxmlElement
    for name, look, outline in (("Interlude Title", "Heading 2", 0), ("Interlude Subtitle", "Normal", None),
                                ("Interlude Lead", "Normal", None), ("Interlude Heading", "Heading 2", None), ("Interlude", "Normal", None)):
        if name not in names:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = doc.styles["Normal"]
            if outline is not None:
                ppr = st.element.get_or_add_pPr()
                ol = OxmlElement("w:outlineLvl")
                ol.set(qn("w:val"), str(outline))
                ppr.append(ol)
        if name not in rpr and look in rpr:
            rpr[name] = rpr[look]
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    n_img = 0
    for p in paras:
        if p.get("image"):
            np_ = doc.add_paragraph()
            np_._element.getparent().replace(np_._element, img_xml[int(p["image"])])
            n_img += 1
            continue
        para = doc.add_paragraph(style=p["style"])
        for text, b, i in parse_markup(p.get("text", "")):
            run = para.add_run()
            if p["style"] in rpr:
                run._element.insert(0, copy.deepcopy(rpr[p["style"]]))
            run.text = text
            run.bold = True if b else None
            run.italic = True if i else None
    doc.core_properties.title = json.load(open(Path(__file__).resolve().parent.parent / "benyoksam" / "book_meta.json", encoding="utf-8"))["title"]
    doc.core_properties.author = "Tuncay Sancak"
    doc.save(a.out)
    chk = Document(a.out).paragraphs
    assert len(chk) == len(paras), (len(chk), len(paras))
    for k, (x, p) in enumerate(zip(chk, paras)):
        if p.get("image"):
            assert x._element.xpath('.//*[local-name()="blip"]'), f"Bild fehlt bei {k}"
        else:
            assert x.text == strip_tags(p["text"]), f"Absatz {k}: Text ≠ Vorlage"
            assert x.style.name == p["style"], f"Absatz {k}: Stil"
    words = sum(len(x.text.split()) for x in chk)
    print(f"{a.out}: {len(chk)} Absätze, {n_img} Bilder, {words} Wörter – OK")


if __name__ == "__main__":
    main()
