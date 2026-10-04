#!/usr/bin/env python3
"""Textänderungen an einem DOCX-Manuskript als Operationsliste anwenden.

Die Operationen stehen in einer JSON-Liste (siehe src/v6_aenderungen.json) und
werden über Text verankert, nicht über Absatznummern – so bleiben sie lesbar und
überstehen vorangehende Einfügungen/Löschungen.

  {"op": "replace", "old": "…", "new": "…", "match": "…"(opt.), "why": "…"}
      ersetzt eine Textstelle innerhalb EINES Absatzes. Formatierung bleibt
      erhalten (die Ersetzung erbt das Format des ersten betroffenen Runs).
      Überdeckt die Stelle Runs mit unterschiedlicher Formatierung (z. B. die
      fetten Akrostichon-Buchstaben in ŞAHİT), wird abgebrochen.
  {"op": "set", "match": "…", "text": "…"}
      ersetzt den ganzen Absatztext (nur bei einheitlich formatierten Absätzen).
  {"op": "delete", "match": "…"}
      löscht den Absatz.
  {"op": "insert_after", "match": "…", "text": "…", "style": "Normal"(opt.)}
      fügt nach dem Absatz einen neuen ein (Absatz- und Zeichenformat vom
      Ankerabsatz, ohne fett/kursiv).
  {"op": "style", "match": "…", "style": "…"}
      setzt die Absatzvorlage.

„match“ (bzw. bei replace „old“) muss genau einen Absatz treffen; mit
"nth": k wird der k-te Treffer (0-basiert) genommen.

Aufruf: python3 tools/docx_ops.py IN.docx OPS.json OUT.docx [--key SA]
"""
import argparse, copy, json, sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE


class OpError(Exception):
    pass


def fmt(r):
    return (bool(r.bold), bool(r.italic), bool(r.underline))


def find_par(doc, needle, nth=None):
    hits = [p for p in doc.paragraphs if needle in p.text]
    if nth is not None:
        if nth >= len(hits):
            raise OpError(f"nur {len(hits)} Treffer für {needle[:60]!r}")
        return hits[nth]
    if len(hits) != 1:
        raise OpError(f"{len(hits)} Treffer statt 1 für {needle[:60]!r}")
    return hits[0]


def replace_in_par(p, old, new):
    text = "".join(r.text for r in p.runs)
    k = text.find(old)
    if k < 0:
        raise OpError(f"Text nicht in den Runs gefunden: {old[:60]!r}")
    if text.count(old) > 1:
        raise OpError(f"Text mehrdeutig im Absatz: {old[:60]!r}")
    start, end = k, k + len(old)
    spans, pos = [], 0
    for r in p.runs:
        spans.append((pos, pos + len(r.text), r))
        pos += len(r.text)
    hit = [(a, b, r) for a, b, r in spans if b > start and a < end]
    if len({fmt(r) for _, _, r in hit}) > 1:
        raise OpError(f"Stelle überdeckt unterschiedlich formatierte Runs: {old[:60]!r}")
    a0, b0, r0 = hit[0]
    tail = r0.text[end - a0:] if end <= b0 else ""
    r0.text = r0.text[:start - a0] + new + tail
    for a, b, r in hit[1:]:
        r.text = r.text[max(0, end - a):] if end < b else ""


def set_text(p, text, force=False):
    runs = [r for r in p.runs if r.text]
    if len({fmt(r) for r in runs}) > 1 and not force:
        raise OpError(f"Absatz gemischt formatiert, 'set' verweigert: {p.text[:60]!r}")
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def ensure_style(doc, name):
    names = {s.name for s in doc.styles}
    if name not in names:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
    return name


def insert_after(doc, p, text, style=None):
    new = copy.deepcopy(p._element)
    p._element.addnext(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, p._parent)
    # alle Runs bis auf den längsten entfernen, Text setzen, fett/kursiv aus
    runs = list(np_.runs)
    keep = max(runs, key=lambda r: len(r.text)) if runs else None
    for r in runs:
        if r is not keep:
            r._element.getparent().remove(r._element)
    if keep is None:
        keep = np_.add_run("")
    keep.text = text
    keep.bold = None
    keep.italic = None
    if style:
        np_.style = doc.styles[ensure_style(doc, style)]
    return np_


def insert_before_heading(doc, p, text):
    """Neue Kapitelüberschrift (Heading 1) vor Absatz p; Format von einer vorhandenen H1."""
    h = next(q for q in doc.paragraphs if q.style.name == "Heading 1")
    new = copy.deepcopy(h._element)
    p._element.addprevious(new)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new, p._parent)
    runs = list(np_.runs)
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    (runs[0] if runs else np_.add_run("")).text = text


def apply(doc, ops, log, locked=frozenset()):
    base = list(doc.paragraphs)          # Ausgangsabsätze, für "pid"
    def target(op):
        if "pid" in op:
            if op["pid"] in locked:
                raise OpError(f"Absatz {op['pid']} ist gesperrt")
            return base[op["pid"]]
        return find_par(doc, op.get("match") or op["old"], op.get("nth"))
    for n, op in enumerate(ops):
        kind = op["op"]
        try:
            if kind == "replace":
                p = target(op)
                replace_in_par(p, op["old"], op["new"])
            elif kind == "set":
                p = target(op)
                set_text(p, op["text"], op.get("force", False))
            elif kind == "split_chapter":
                p = target(op)
                insert_before_heading(doc, p, op["title"])
                if "Opening" in {st.name for st in doc.styles}:
                    p.style = doc.styles["Opening"]
            elif kind == "delete":
                p = target(op)
                if p.style.name.startswith("Heading") or p.style.name in ("Title", "Book Part"):
                    raise OpError("Überschriften werden nicht gelöscht")
                if any(r.bold and len(r.text.strip()) <= 3 and r.text.strip() for r in p.runs):
                    raise OpError("Absatz trägt Akrostichon-Fettbuchstaben")
                p._element.getparent().remove(p._element)
            elif kind == "insert_after":
                p = target(op)
                insert_after(doc, p, op["text"], op.get("style"))
            elif kind == "delete_empty":
                # leere Absätze zwischen zwei Ankern (ausschließlich) löschen
                ps = doc.paragraphs
                els = [q._element for q in ps]
                a = els.index(find_par(doc, op["from"])._element) if op.get("from") else -1
                b = els.index(find_par(doc, op["to"])._element)
                for q in ps[a + 1:b]:
                    if not q.text.strip() and not q._element.xpath(".//w:drawing"):
                        q._element.getparent().remove(q._element)
            elif kind == "strip_ws_runs":
                # reine Leerraum-Runs (Tabs) am Absatzanfang entfernen
                p = find_par(doc, op["match"], op.get("nth"))
                for r in list(p.runs):
                    if r.text.strip():
                        break
                    r._element.getparent().remove(r._element)
            elif kind == "clean_whitespace":
                # führende/abschließende Leerzeichen und Tabs, doppelte Leerzeichen
                import re as _re
                for q in doc.paragraphs:
                    rs = [r for r in q.runs if r.text]
                    for r in rs:
                        if "  " in r.text:
                            r.text = _re.sub(r" {2,}", " ", r.text)
                    for r in rs:
                        t = r.text.lstrip(" \t\u00a0")
                        r.text = t
                        if t:
                            break
                    for r in reversed(rs):
                        t = r.text.rstrip(" \t\u00a0")
                        r.text = t
                        if t:
                            break
            elif kind == "regex_all":
                # Muster innerhalb einzelner Runs ersetzen, in Absätzen zwischen zwei Ankern (einschließlich)
                import re as _re
                ps = doc.paragraphs
                els = [q._element for q in ps]
                a = els.index(find_par(doc, op["from"])._element)
                b = els.index(find_par(doc, op["to"])._element)
                n = 0
                for q in ps[a:b + 1]:
                    for r in q.runs:
                        t2, k = _re.subn(op["pattern"], op["repl"], r.text)
                        if k:
                            r.text = t2; n += k
                if n != op.get("count", n):
                    raise OpError(f"{n} Ersetzungen statt {op['count']}")
            elif kind == "style":
                p = find_par(doc, op["match"], op.get("nth"))
                p.style = doc.styles[ensure_style(doc, op["style"])]
            else:
                raise OpError(f"unbekannte Operation {kind!r}")
            log.append({"n": n, "ok": True})
        except OpError as e:
            log.append({"n": n, "ok": False, "fehler": str(e), "op": op})
    return log


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inp"); ap.add_argument("ops"); ap.add_argument("out")
    ap.add_argument("--key", help="Schlüssel in der JSON-Datei (z. B. SA/BY)")
    ap.add_argument("--allow-errors", action="store_true")
    ap.add_argument("--locked", help="JSON-Liste gesperrter Absatz-IDs (pid)")
    a = ap.parse_args()
    data = json.load(open(a.ops))
    ops = data[a.key] if a.key else data
    doc = Document(a.inp)
    locked = frozenset(json.load(open(a.locked))) if a.locked else frozenset()
    log = apply(doc, ops, [], locked)
    bad = [l for l in log if not l["ok"]]
    for b in bad:
        print("FEHLER", b["n"], b["fehler"], json.dumps(b["op"], ensure_ascii=False)[:300], file=sys.stderr)
    if bad and not a.allow_errors:
        sys.exit(1)
    doc.save(a.out)
    print(f"{len(log) - len(bad)}/{len(log)} Operationen angewandt → {a.out}")


if __name__ == "__main__":
    main()
