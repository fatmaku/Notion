#!/usr/bin/env python3
"""Änderungsliste (Markdown, Türkisch) aus Operationslisten erzeugen.

Aufruf: python3 tools/changelist.py BASE.docx NEU.docx OUT.md --title "…" OPS1.json [OPS2.json …]
Je Operation: Kapitel, Art (silindi / kısaltıldı / değiştirildi / eklendi / yeni bölüm),
Ausschnitt alt → neu, Begründung („neden“/„why“). Vorweg eine Tabelle Wörter je Kapitel.
"""
import argparse, json
from docx import Document


def chapters_of(paras):
    cur, out = "(Ön kısım)", []
    for p in paras:
        if p.style.name == "Heading 1":
            cur = p.text.strip()
        out.append(cur)
    return out


def words_per_chapter(paras):
    rows, cur = {}, "(Ön kısım)"
    order = [cur]
    for p in paras:
        if p.style.name == "Heading 1":
            cur = p.text.strip(); order.append(cur); rows.setdefault(cur, 0)
            continue
        rows[cur] = rows.get(cur, 0) + len(p.text.split())
    return order, rows


def short(s, n=160):
    s = " ".join(s.split())
    return s if len(s) <= n else s[: n - 1] + "…"


def md_escape(s):
    return s.replace("|", "\\|")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("base"); ap.add_argument("new"); ap.add_argument("out")
    ap.add_argument("--title", default="Değişiklik listesi")
    ap.add_argument("--intro", default="")
    ap.add_argument("ops", nargs="+")
    a = ap.parse_args()
    bp = Document(a.base).paragraphs
    np_ = Document(a.new).paragraphs
    chap = chapters_of(bp)
    o1, w1 = words_per_chapter(bp)
    o2, w2 = words_per_chapter(np_)
    L = [f"# {a.title}", ""]
    if a.intro:
        L += [a.intro, ""]
    t1, t2 = sum(w1.values()), sum(w2.values())
    L += [f"**Toplam:** {t1:,} → {t2:,} kelime ({t2 - t1:+,}; {100 * (t2 - t1) / t1:+.1f} %)".replace(",", "."), "",
          "| Bölüm | Önce | Sonra | Fark |", "|---|---:|---:|---:|"]
    seen = set()
    for name in o2:
        b, n = w1.get(name, 0), w2.get(name, 0)
        seen.add(name)
        L.append(f"| {md_escape(name)} | {b} | {n} | {n - b:+d} |")
    for name in o1:
        if name not in seen:
            L.append(f"| {md_escape(name)} | {w1[name]} | – | {-w1[name]:+d} |")
    L.append("")
    n_total = 0
    for path in a.ops:
        ops = json.load(open(path))
        name = path.split("/")[-1]
        HEAD = {"ops_sa_entegrasyon.json": "Edisyon'un işlenmesi (başlık, ithaf, köprüler, üç roman, yazım)",
                "ops_sa_noktalama_ek.json": "Ek noktalama ve boşluk temizliği",
                "ops_SA_alle.json": "Sıkılaştırma (beş lektör bloğu)", "ops_BY_alle.json": "Sıkılaştırma (altı lektör bloğu)",
                "ops_editor_SA.json": "Editör müdahaleleri (jüri önerileri, lektör notları)",
                "ops_editor_BY.json": "Editör müdahaleleri (jüri önerileri, lektör notları)",
                "fix_SA_alle.json": "Bağımsız kontrol sonrası dikiş düzeltmeleri", "fix_BY_alle.json": "Bağımsız kontrol sonrası dikiş düzeltmeleri",
                "kor_SA_alle.json": "Son okuma (düzelti): yazım, dilbilgisi, noktalama, tutarlılık", "kor_BY_alle.json": "Son okuma (düzelti): yazım, dilbilgisi, noktalama, tutarlılık",
                "ops_editor2_SA.json": "Jüri v6 sonrası editör müdahaleleri (içerik tutarlılığı, köprüler)", "ops_editor2_BY.json": "Jüri v6 sonrası editör müdahaleleri (içerik tutarlılığı)"}
        L += [f"## {HEAD.get(name, name)}", ""]
        cur = None
        for op in ops:
            if op.get("op") in ("delete_empty", "clean_whitespace", "strip_ws_runs") or op.get("_skip"):
                continue
            kind = op["op"]
            pid = op.get("pid")
            kap = op.get("kap") or (chap[pid] if pid is not None and pid < len(chap) else "")
            why = op.get("neden") or op.get("why") or op.get("sorun") or op.get("not") or ""
            if op.get("tur"):
                why = f"[{op['tur']}] {why}"
            if kap != cur:
                L += ["", f"### {kap}", ""]
                cur = kap
            base_text = bp[pid].text if pid is not None and pid < len(bp) else ""
            if kind == "delete":
                src = base_text or op.get("match", "")
                L.append(f"- **silindi:** „{short(src)}“ — {why}")
            elif kind == "replace":
                old, new = op["old"], op["new"]
                tag = "kısaltıldı" if not new.strip() or len(new) < len(old) * 0.6 else "değiştirildi"
                if new.strip():
                    L.append(f"- **{tag}:** „{short(old)}“ → „{short(new)}“ — {why}")
                else:
                    L.append(f"- **{tag}:** „{short(old)}“ çıkarıldı — {why}")
            elif kind == "set":
                L.append(f"- **yeniden kuruldu:** „{short(base_text)}“ → „{short(op['text'])}“ — {why}")
            elif kind == "insert_after":
                L.append(f"- **eklendi:** „{short(op['text'], 400)}“ — {why}")
            elif kind == "split_chapter":
                L.append(f"- **yeni bölüm:** „{op['title']}“ bu paragrafla başlıyor: „{short(base_text, 100)}“ — {why}")
            elif kind == "style":
                L.append(f"- **biçim:** „{short(op.get('match', ''), 80)}“ → {op['style']} — {why}")
            n_total += 1
        L.append("")
    L.insert(2, f"Toplam {n_total} işlem.\n")
    open(a.out, "w").write("\n".join(L) + "\n")
    print(a.out, n_total, "Einträge")


if __name__ == "__main__":
    main()
