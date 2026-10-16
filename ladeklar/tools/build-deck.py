#!/usr/bin/env python3
"""Pitch deck as editable PowerPoint (16:9) in the Ladeklar design, content from business/PITCH.md.
   python3 tools/build-deck.py  ->  business/Ladeklar_Pitch.pptx (+ PDF via LibreOffice)
Fonts: Archivo + Manrope (business/fonts, SIL OFL) – install them for the exact look."""
import re, os, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = lambda f: os.path.join(ROOT, 'business', 'deck-assets', f)
INK, INK2, BRASS, BRASS2 = RGBColor(0x0B, 0x16, 0x22), RGBColor(0x13, 0x23, 0x3A), RGBColor(0xC8, 0xA4, 0x5C), RGBColor(0xE6, 0xCD, 0x92)
WHITE, MUTED, LINE, TEAL, OK, BAD = RGBColor(0xF2, 0xF5, 0xF8), RGBColor(0x9F, 0xB0, 0xC3), RGBColor(0x2A, 0x3C, 0x50), RGBColor(0x0E, 0x7C, 0x7B), RGBColor(0x3F, 0xC4, 0x83), RGBColor(0xF0, 0x72, 0x72)
H, B = 'Archivo ExtraBold', 'Manrope'

# ---------- parse PITCH.md ----------
md = open(os.path.join(ROOT, 'business', 'PITCH.md'), encoding='utf-8').read()
slides = []
for block in re.split(r'\n## Folie \d+ – ', md)[1:]:
    name = block.split('\n', 1)[0].strip()
    head = re.search(r'\*\*Überschrift:\*\*\s*(.+)', block).group(1).strip()
    big = re.search(r'\*\*Große Kennzahl:\*\*\s*(.+)', block)
    big = big.group(1).strip() if big else ''
    bullets = []
    sec = block.split('**Große Kennzahl:**')[0]
    for l in sec.split('\n'):
        if l.startswith('- '):
            bullets.append(l[2:].strip())
    notes = re.search(r'\*\*Sprechertext[^*]*\*\*\s*(.+?)(?:\n\n(?:Quellen|---)|\Z)', block, re.S)
    notes = re.sub(r'\s+', ' ', notes.group(1)).strip() if notes else ''
    regie = re.search(r'\*\*Regiehinweise[^*]*\*\*\s*(.+?)\n\n\*\*', block, re.S)
    if regie:
        notes = 'REGIE: ' + re.sub(r'\s+', ' ', regie.group(1)).strip() + '\n\n' + notes
    table = re.findall(r'^\|.+\|$', sec, re.M)
    slides.append(dict(name=name, head=head, big=big, bullets=bullets, notes=notes, table=table))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def clean(s):
    return re.sub(r'\[Q\d+\](\[Q\d+\])*', '', s).replace('**', '').strip()


def text(slide, x, y, w, h, s, size=16, color=WHITE, font=B, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.12):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = line if isinstance(line, list) else [(line, color)]
        for seg, col in parts:
            r = p.add_run()
            r.text = seg
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = col
    return tb


def bullets(slide, x, y, w, items, size=15, gap=10, color=WHITE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = '—  '
        r.font.size = Pt(size)
        r.font.name = B
        r.font.bold = True
        r.font.color.rgb = BRASS
        m = re.match(r'\*\*(.+?)\*\*\s*(.*)', item)
        if m:
            r2 = p.add_run(); r2.text = clean(m.group(1)) + ' '; r2.font.size = Pt(size); r2.font.name = B; r2.font.bold = True; r2.font.color.rgb = WHITE
            rest = m.group(2)
        else:
            rest = item
        r3 = p.add_run(); r3.text = clean(rest); r3.font.size = Pt(size); r3.font.name = B; r3.font.color.rgb = color
    return tb


def box(slide, x, y, w, h, fill=INK2, line=None, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def picture(slide, f, x, y, w=None, h=None, rot=0):
    pic = slide.shapes.add_picture(A(f), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    pic.rotation = rot
    return pic


def phone(slide, f, x, y, h, rot=0):
    """Image of a phone screen inside a rounded dark device frame."""
    w = h * 1170 / 2532
    frame = box(slide, x - 0.09, y - 0.09, w + 0.18, h + 0.18, fill=RGBColor(0x1A, 0x22, 0x2D), line=RGBColor(0x3C, 0x48, 0x59), radius=0.16)
    frame.rotation = rot
    pic = picture(slide, f, x, y, w, h, rot)
    return pic


def transparent(chart):
    from pptx.oxml.ns import qn
    from lxml import etree
    cs = chart._chartSpace
    for parent in (cs, cs.find(qn('c:chart')).find(qn('c:plotArea'))):
        sp = parent.find(qn('c:spPr'))
        if sp is None:
            sp = etree.SubElement(parent, qn('c:spPr'))
        for child in list(sp):
            sp.remove(child)
        etree.SubElement(sp, qn('a:noFill'))
        ln = etree.SubElement(sp, qn('a:ln')); etree.SubElement(ln, qn('a:noFill'))
    # c:spPr must come before c:txPr / c:externalData etc. in chartSpace – move it after c:chart
    sp = cs.find(qn('c:spPr')); cs.remove(sp); cs.find(qn('c:chart')).addnext(sp)


def frame(slide, n, total, variant='a'):
    slide.shapes.add_picture(A(f'bg-{variant}.png'), 0, 0, prs.slide_width, prs.slide_height)
    text(slide, 0.6, 0.42, 4, 0.3, 'LADEKLAR', 11, BRASS2, H)
    text(slide, 9.7, 7.0, 3.05, 0.3, f'{n:02d} / {total:02d}', 10, MUTED, B, align=PP_ALIGN.RIGHT)
    text(slide, 0.6, 7.0, 6, 0.3, 'Ein Tipp. Alles klar.', 10, MUTED, B)


def headline(slide, s, y=1.0, w=7.2, size=36):
    return text(slide, 0.6, y, w, 1.6, clean(s), size, WHITE, H, spacing=1.0)


def bignum(slide, s, x, y, w=4.6, size=54, label=''):
    m = re.match(r'\*\*(.+?)\*\*\s*(.*)', s)
    num, rest = (m.group(1), m.group(2)) if m else (s, '')
    text(slide, x, y, w, 1.0, clean(num), size, BRASS2, H, spacing=0.95)
    if rest or label:
        text(slide, x, y + size / 72 * 1.15 + 0.05, w, 0.8, clean(label or rest), 12.5, MUTED, B)


N = len(slides)
for i, s in enumerate(slides, 1):
    sl = prs.slides.add_slide(BLANK)
    sl.notes_slide.notes_text_frame.text = s['notes']
    nm = s['name']
    if nm.startswith('Titel'):
        frame(sl, i, N, 'a')
        text(sl, 0.6, 1.55, 7, 0.4, 'PITCH · OKTOBER 2026', 12, BRASS2, B, bold=True)
        text(sl, 0.6, 2.0, 7.6, 2.4, [[('Ladeklar.', WHITE)], [('Ein Tipp. ', WHITE), ('Alles klar.', BRASS2)]], 54, WHITE, H, spacing=0.95)
        bullets(sl, 0.6, 4.35, 6.8, s['bullets'], 15)
        phone(sl, 'phone-2-loaded.png', 9.35, 0.85, 5.9, rot=0)
    elif nm.startswith('Problem'):
        frame(sl, i, N, 'b')
        headline(sl, s['head'], w=6.5, size=34)
        bullets(sl, 0.6, 2.25, 6.3, s['bullets'], 14.5)
        bignum(sl, s['big'], 0.6, 5.2, 6.2, 44)
        picture(sl, 'chaos-chat.jpg', 7.35, 1.05, 5.4)
    elif nm.startswith('Lösung'):
        frame(sl, i, N, 'a')
        headline(sl, s['head'])
        bullets(sl, 0.6, 2.05, 6.6, s['bullets'], 15)
        bignum(sl, s['big'], 0.6, 5.55, 5, 46, 'Der Fahrer sieht immer nur den nächsten Schritt.')
        phone(sl, 'phone-1b-ok.png', 8.0, 1.15, 5.2, rot=-4)
        phone(sl, 'phone-2-loaded.png', 10.35, 0.75, 5.6, rot=3)
    elif nm.startswith('So funktioniert'):
        frame(sl, i, N, 'b')
        headline(sl, s['head'], w=11)
        for k, b in enumerate(s['bullets'][:3]):
            m = re.match(r'\*\*(\d)\.\s*(.+?):?\*\*:?\s*(.*)', b)
            num, title, body = (m.group(1), m.group(2).rstrip(':'), m.group(3)) if m else (str(k + 1), b, '')
            x = 0.6 + k * 4.15
            box(sl, x, 2.35, 3.85, 3.55, fill=INK2, line=LINE)
            text(sl, x + 0.35, 2.6, 1.5, 1.2, num, 60, BRASS, H)
            text(sl, x + 0.35, 3.85, 3.2, 0.6, title, 19, WHITE, H)
            text(sl, x + 0.35, 4.45, 3.2, 1.4, clean(body), 13.5, MUTED, B, spacing=1.2)
        bignum(sl, s['big'], 0.6, 6.05, 8, 26)
    elif nm.startswith('Live-Demo'):
        frame(sl, i, N, 'c')
        headline(sl, s['head'], w=11, size=32)
        picture(sl, 'demo-split.png', 0.6, 1.85, 8.3)
        bullets(sl, 9.25, 2.0, 3.6, s['bullets'], 12.5, 8)
        bignum(sl, s['big'], 9.25, 5.35, 3.6, 40, 'Live zwischen Handy und Laptop')
    elif nm.startswith('Markt'):
        frame(sl, i, N, 'a')
        headline(sl, s['head'], w=7.5)
        bullets(sl, 0.6, 2.2, 6.6, s['bullets'], 13.5, 8)
        bignum(sl, s['big'], 7.9, 1.25, 4.8, 54)
        for k, (lab, val, w) in enumerate([('TAM Deutschland', '≈ 140 Mio. €', 4.6), ('SAM erreichbar', '≈ 92 Mio. €', 3.05), ('Ziel Jahr 5 (ARR)', '≈ 2,0 Mio. €', 0.55)]):
            y = 3.35 + k * 1.05
            box(sl, 7.9, y, w, 0.62, fill=[RGBColor(0x2A, 0x3C, 0x50), TEAL, BRASS][k], radius=0.25)
            text(sl, 7.9, y + 0.68, 4.8, 0.3, f'{lab}  ·  {val}', 12, WHITE if k < 2 else BRASS2, B, bold=True)
    elif nm.startswith('Geschäftsmodell'):
        frame(sl, i, N, 'b')
        headline(sl, s['head'], w=11)
        plans = [('Start', '49 €', '3 Fahrer · 6 Fahrzeuge'), ('Flotte', '129 €', '15 Fahrer · +7 € je weiterer'), ('Premium', '299 €', '40 Fahrer · API · SLA 99,5 %')]
        for k, (n, p, d) in enumerate(plans):
            x = 0.6 + k * 3.0
            hl = k == 1
            box(sl, x, 2.2, 2.75, 2.55, fill=INK2, line=BRASS if hl else LINE)
            text(sl, x + 0.3, 2.45, 2.2, 0.3, n.upper(), 11, BRASS2 if hl else MUTED, B, bold=True)
            text(sl, x + 0.3, 2.8, 2.3, 0.9, p, 40, WHITE, H)
            text(sl, x + 0.3, 3.65, 2.2, 0.3, 'netto / Monat', 11, MUTED, B)
            text(sl, x + 0.3, 4.0, 2.25, 0.7, d, 12, WHITE, B)
        bullets(sl, 0.6, 5.05, 8.6, s['bullets'][1:], 12.5, 6)
        bignum(sl, s['big'], 9.75, 2.2, 3.1, 50)
    elif nm.startswith('Wettbewerb'):
        frame(sl, i, N, 'a')
        headline(sl, s['head'], w=6.4, size=32)
        bullets(sl, 0.6, 2.1, 6.0, s['bullets'], 12.5, 7)
        # positioning matrix
        X0, Y0, W, Hh = 7.2, 1.25, 5.4, 4.6
        box(sl, X0, Y0, W, Hh, fill=INK2, line=LINE, radius=0.04)
        ln = sl.shapes.add_connector(1, Inches(X0 + W / 2), Inches(Y0 + 0.2), Inches(X0 + W / 2), Inches(Y0 + Hh - 0.2)); ln.line.color.rgb = LINE
        ln = sl.shapes.add_connector(1, Inches(X0 + 0.2), Inches(Y0 + Hh / 2), Inches(X0 + W - 0.2), Inches(Y0 + Hh / 2)); ln.line.color.rgb = LINE
        text(sl, X0, Y0 + Hh + 0.08, W, 0.3, '←  sofort startklar          Einführungsaufwand          Projekt / Hardware  →', 9.5, MUTED, B, align=PP_ALIGN.CENTER)
        text(sl, X0 - 0.05, Y0 - 0.32, W, 0.3, '↑ Fahrer & Disposition im Mittelpunkt', 9.5, MUTED, B)
        pts = [('Papier · WhatsApp', 0.12, 0.62, MUTED), ('App-Tools', 0.42, 0.3, MUTED), ('Telematik', 0.78, 0.7, MUTED), ('TMS', 0.86, 0.35, MUTED), ('Ladeklar', 0.14, 0.12, BRASS)]
        for lab, px, py, col in pts:
            r = 0.34 if lab == 'Ladeklar' else 0.22
            c = box(sl, X0 + px * W - r / 2, Y0 + py * Hh - r / 2, r, r, fill=col, shape=MSO_SHAPE.OVAL)
            text(sl, X0 + px * W + 0.25, Y0 + py * Hh - 0.15, 2.2, 0.3, lab, 13 if lab == 'Ladeklar' else 11, BRASS2 if lab == 'Ladeklar' else WHITE, B, bold=lab == 'Ladeklar')
        bignum(sl, s['big'], 0.6, 5.75, 6.4, 30)
    elif nm.startswith('Traktion'):
        frame(sl, i, N, 'b')
        headline(sl, s['head'], w=7.4, size=30)
        for k, b in enumerate(s['bullets']):
            y = 2.45 + k * 0.8
            c = box(sl, 0.6, y + 0.04, 0.34, 0.34, fill=OK if k < 2 else LINE, shape=MSO_SHAPE.OVAL)
            text(sl, 0.62, y + 0.06, 0.3, 0.3, '✓' if k < 2 else str(k - 1), 12, INK if k < 2 else WHITE, B, bold=True, align=PP_ALIGN.CENTER)
            text(sl, 1.15, y, 6.9, 0.75, clean(b), 14, WHITE, B)
        bignum(sl, s['big'], 8.4, 1.15, 4.6, 36)
        phone(sl, 'phone-5-pod.png', 10.1, 2.55, 4.15)
    elif nm.startswith('Go-to-Market'):
        frame(sl, i, N, 'a')
        headline(sl, s['head'], w=8.5, size=32)
        steps = [('Q1 2027', 'Gründer-Vertrieb, Google Ads, Postkarten'), ('März 2027', 'LogiMAT Stuttgart'), ('April 2027', 'transport logistic München'), ('Juli 2027', 'Vertrieb & Kundenerfolg an Bord'), ('2028', 'Partner: Verbände, Telematik, Steuerberater')]
        ln = sl.shapes.add_connector(1, Inches(0.8), Inches(2.62), Inches(12.6), Inches(2.62)); ln.line.color.rgb = BRASS; ln.line.width = Pt(2)
        for k, (when, what) in enumerate(steps):
            x = 0.6 + k * 2.45
            box(sl, x + 0.12, 2.5, 0.24, 0.24, fill=BRASS, shape=MSO_SHAPE.OVAL)
            text(sl, x, 2.0, 2.3, 0.35, when, 13, BRASS2, B, bold=True)
            text(sl, x, 2.95, 2.25, 1.0, what, 13, WHITE, B)
        bullets(sl, 0.6, 4.25, 7.6, s['bullets'][2:], 12.5, 6)
        bignum(sl, s['big'], 8.9, 4.2, 4, 46)
    elif nm.startswith('Finanzen'):
        frame(sl, i, N, 'b')
        headline(sl, s['head'], w=8)
        rows = [[c.strip() for c in r.strip('|').split('|')] for r in s['table'] if not re.match(r'^\|[-| ]+\|$', r)]
        tbl = sl.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(2.1), Inches(6.9), Inches(2.3)).table
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = INK if r == 0 else INK2
                cell.text = val
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
                for run in p.runs:
                    run.font.size = Pt(11.5)
                    run.font.name = B
                    run.font.bold = r == 0 or c == 0
                    run.font.color.rgb = BRASS2 if r == 0 else (BAD if val.startswith('−') else WHITE)
        bullets(sl, 0.6, 4.75, 6.9, s['bullets'], 12.5, 6)
        cd = CategoryChartData()
        cd.categories = ['2027', '2028', '2029', '2030', '2031']
        cd.add_series('ARR (Mio. €)', (0.13, 0.44, 0.90, 1.41, 1.98))
        ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.9), Inches(1.9), Inches(4.9), Inches(3.9), cd).chart
        ch.has_legend = False
        ch.has_title = False
        transparent(ch)
        ch.plots[0].gap_width = 60
        ser = ch.plots[0].series[0]
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = BRASS
        ch.plots[0].has_data_labels = True
        dl = ch.plots[0].data_labels; dl.number_format = '0.00'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(11); dl.font.color.rgb = WHITE; dl.font.name = B
        ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
        ch.category_axis.tick_labels.font.color.rgb = MUTED; ch.category_axis.tick_labels.font.size = Pt(11); ch.category_axis.format.line.color.rgb = LINE
        text(sl, 7.9, 1.6, 4.9, 0.3, 'Wiederkehrender Jahresumsatz (ARR, Mio. €)', 11, MUTED, B)
        bignum(sl, s['big'], 7.9, 5.95, 4.9, 26)
    elif nm.startswith('Team'):
        frame(sl, i, N, 'a')
        headline(sl, s['head'], w=7)
        bullets(sl, 0.6, 2.1, 6.5, s['bullets'], 13, 7)
        box(sl, 8.55, 1.35, 3.4, 0.9, fill=BRASS)
        text(sl, 8.55, 1.42, 3.4, 0.8, ['LIZENZGEBER', 'Plattform · Lizenzen · Abrechnung'], 11.5, INK, B, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        for k, (n, col) in enumerate([('Kaya Logistik', TEAL), ('Nordlicht', RGBColor(0x25, 0x63, 0xEB)), ('Ihre Kunden …', LINE)]):
            x = 7.55 + k * 1.85
            ln = sl.shapes.add_connector(1, Inches(10.25), Inches(2.25), Inches(x + 0.8), Inches(3.05)); ln.line.color.rgb = BRASS
            box(sl, x, 3.05, 1.65, 0.85, fill=col)
            text(sl, x, 3.1, 1.65, 0.75, [n, 'eigenes Logo'], 10.5, WHITE, B, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            for d in range(3):
                box(sl, x + 0.18 + d * 0.48, 4.25, 0.36, 0.36, fill=INK2, line=LINE, shape=MSO_SHAPE.OVAL)
        text(sl, 7.55, 4.75, 5.4, 0.4, 'Fahrer je Kunde – mit PIN, in 8 Sprachen', 11, MUTED, B, align=PP_ALIGN.CENTER)
        bignum(sl, s['big'], 7.55, 5.35, 5.4, 34)
    elif nm == 'Ask':
        frame(sl, i, N, 'c')
        text(sl, 0.6, 1.0, 7, 0.4, 'ASK', 12, BRASS2, B, bold=True)
        text(sl, 0.6, 1.35, 7.5, 1.3, '200.000 €', 72, WHITE, H)
        text(sl, 0.6, 2.65, 7.2, 0.6, clean(s['head']), 20, BRASS2, H)
        bullets(sl, 0.6, 3.45, 7.0, s['bullets'], 13, 7)
        cd = CategoryChartData()
        cd.categories = ['Personal', 'Marketing & Messen', 'Betrieb', 'Variable Kosten']
        cd.add_series('Verwendung', (0.64, 0.18, 0.13, 0.05))
        ch = sl.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(8.4), Inches(1.3), Inches(4.4), Inches(4.6), cd).chart
        transparent(ch)
        ch.has_title = False
        ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
        ch.legend.font.color.rgb = WHITE; ch.legend.font.size = Pt(10.5); ch.legend.font.name = B
        for k, col in enumerate([BRASS, TEAL, RGBColor(0x5B, 0x6E, 0x86), RGBColor(0x2A, 0x3C, 0x50)]):
            pt = ch.plots[0].series[0].points[k]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = col
        ch.plots[0].has_data_labels = True
        dl = ch.plots[0].data_labels; dl.number_format = '0%'; dl.number_format_is_linked = False; dl.font.size = Pt(11); dl.font.color.rgb = WHITE; dl.font.bold = True
    elif nm.startswith('Kontakt'):
        frame(sl, i, N, 'c')
        text(sl, 0, 2.0, 13.333, 1.3, [[('Ein Tipp. ', WHITE), ('Alles klar.', BRASS2)]], 60, WHITE, H, align=PP_ALIGN.CENTER)
        text(sl, 0, 3.45, 13.333, 0.5, '[Name] · Gründer und Lizenzgeber', 18, WHITE, B, bold=True, align=PP_ALIGN.CENTER)
        text(sl, 0, 3.95, 13.333, 0.5, '[E-Mail] · [Telefon] · [Website]', 16, MUTED, B, align=PP_ALIGN.CENTER)
        b = box(sl, 5.17, 4.75, 3.0, 0.7, fill=BRASS, radius=0.3)
        text(sl, 5.17, 4.83, 3.0, 0.55, '30 Tage kostenlos testen', 15, INK, B, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        frame(sl, i, N, 'a')
        headline(sl, s['head'])
        bullets(sl, 0.6, 2.2, 7, s['bullets'])

out = os.path.join(ROOT, 'business', 'Ladeklar_Pitch.pptx')
prs.save(out)
print('saved', out, len(prs.slides), 'Folien')
pdfdir = os.path.join(ROOT, 'business', 'pdf')
os.makedirs(pdfdir, exist_ok=True)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', pdfdir, out], check=True, capture_output=True)
print('pdf', os.path.join(pdfdir, 'Ladeklar_Pitch.pdf'))
