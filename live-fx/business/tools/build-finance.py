#!/usr/bin/env python3
"""
LiveFX – Finanzmodell und Bewertung für die 250.000-€-Pre-Seed-Runde

    python3 tools/build-finance.py            # baut LiveFX_Finanzmodell.xlsx + tools/finance-250k.json
    python3 tools/build-finance.py --no-recalc  # nur die Excel-Datei schreiben (ohne Werte/JSON)

Ablauf:
  1. Schreibt die Arbeitsmappe mit echten Excel-Formeln (openpyxl). Alle Eingaben stehen blau
     auf dem Blatt „Annahmen“ (bzw. Bewertungs-Eingaben auf „Bewertung“), alles andere rechnet.
  2. Lässt LibreOffice alle Formeln berechnen (recalc.py des xlsx-Skills, Pfad über
     $LIVEFX_RECALC; sonst Roundtrip `soffice --headless --convert-to xlsx`).
  3. Prüft alle Zellen auf Fehlerwerte (#REF!, #DIV/0!, #NAME?, #VALUE!, #N/A …) und bricht ab,
     wenn einer gefunden wird.
  4. Liest die berechneten Werte und schreibt tools/finance-250k.json (alle Zahlen für die
     Dokumente: Mittelverwendung, Runway, Kasse M6/12/18/24, Folgerunde, GuV 5 Jahre,
     Bewertung, Gates, Sensitivität, Quellen).

Abhängigkeiten: Python 3, openpyxl, LibreOffice (soffice).
Alle Zahlen sind Schätzungen bzw. ein Vorschlag – mit Steuer-/Rechtsberater prüfen.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
BUSINESS = os.path.dirname(HERE)
XLSX = os.path.join(BUSINESS, "LiveFX_Finanzmodell.xlsx")
JSON_OUT = os.path.join(HERE, "finance-250k.json")
DEFAULT_RECALC = ("/root/.claude/skills/synced/d0dcd7ae-a634-44e6-9a7e-3c377de3ce05_"
                  "22cf4469-625c-4938-a4f2-7b77f06ca164/xlsx/scripts/recalc.py")

# ---------------------------------------------------------------- Stil
FONT = "Arial"
F_IN = Font(name=FONT, color="0000FF")              # Eingabe (blau)
F_CALC = Font(name=FONT, color="000000")            # Formel (schwarz)
F_LINK = Font(name=FONT, color="008000")            # Verweis auf anderes Blatt (grün)
F_BOLD = Font(name=FONT, bold=True)
F_TITLE = Font(name=FONT, bold=True, size=14)
F_H = Font(name=FONT, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, italic=True, color="555555", size=9)
FILL_H = PatternFill("solid", fgColor="1F3A5F")
FILL_SUB = PatternFill("solid", fgColor="DCE6F1")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_TOT = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
B_TOP = Border(top=Side(style="thin", color="000000"))

EUR = '#,##0" €";(#,##0" €");"–"'
TEUR = '#,##0.0" T€";(#,##0.0" T€");"–"'
NUM = '#,##0;(#,##0);"–"'
PCT = '0.0%;(0.0%);"–"'
MULT = '0.0"x"'
MON = '0" M."'

SCEN = ["kons", "basis", "opt"]
SCEN_DE = {"kons": "Konservativ", "basis": "Basis", "opt": "Optimistisch"}
YEARS = ["J1 2027", "J2 2028", "J3 2029", "J4 2030", "J5 2031"]

R = {}  # Registry: Schlüssel -> (Blatt, Zelle) für JSON-Export und Querverweise


def ref(key, absolute=True):
    sh, cell = R[key]
    if absolute:
        col = "".join(c for c in cell if c.isalpha())
        row = "".join(c for c in cell if c.isdigit())
        cell = f"${col}${row}"
    return f"'{sh}'!{cell}"


def reg(key, ws, cell):
    R[key] = (ws.title, cell)


def put(ws, cell, value, font=None, fmt=None, fill=None, bold=False, key=None, comment=None,
        align=None, wrap=False):
    c = ws[cell]
    c.value = value
    if font is None:
        font = F_CALC
        if isinstance(value, str) and value.startswith("="):
            font = F_LINK if "!" in value else F_CALC
    if bold:
        font = Font(name=font.name, bold=True, color=font.color, size=font.size)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if key:
        reg(key, ws, cell)
    if comment:
        c.comment = Comment(comment, "LiveFX-Modell")
    if align or wrap:
        c.alignment = Alignment(horizontal=align, wrap_text=wrap, vertical="top")
    return c


def header_row(ws, row, labels, start_col=1, fill=FILL_H, font=F_H):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = font
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def section(ws, row, text, ncols=8):
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = FILL_SUB
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_BOLD


def note(ws, cell, text):
    c = ws[cell]
    c.value = text
    c.font = F_NOTE
    c.alignment = Alignment(wrap_text=False, vertical="top")


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


DISCLAIMER = ("Alle Werte sind Schätzungen auf Basis von BUSINESSPLAN.md §11 – es gibt noch keine Umsätze. "
              "Mittelverwendung, Gates und Bewertung: Vorschlag – mit Steuer-/Rechtsberater prüfen.")

# =====================================================================================
wb = Workbook()

# ------------------------------------------------------------------ Blatt: Annahmen
A = wb.active
A.title = "Annahmen"
widths(A, {"A": 58, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 70})
put(A, "A1", "LiveFX – Finanzmodell Pre-Seed 250.000 € · Annahmen", F_TITLE)
note(A, "A2", DISCLAIMER)
note(A, "A3", "Legende: blaue Schrift = Eingabe (änderbar) · schwarz = Formel · grün = Verweis auf anderes Blatt · "
              "gelb = Schlüsselannahme. Stand: Oktober 2026, Produkt v2.1. J1 = 2027 (Start Pro-Abo, Closing Jan 2027).")

r = 5
section(A, r, "A. Allgemein", 9)
gen = [
    ("start_cash", "Startkasse (Pre-Seed, Closing Monat 1 = Jan 2027)", 250000, EUR, True,
     "Vom Gründer korrigierte Rundengröße: ca. 250.000 € statt 500.000 € (BP §11.6 alt)."),
    ("infra_fix", "Infrastruktur fix (Server, Domains, Tools) €/Monat", 250, EUR, False,
     "Schätzung; BP §11.3 Infrastruktur J1 8 T€ – wird hier als fix + variabel modelliert."),
    ("infra_var", "Infrastruktur variabel (KI-API nur Pro, Zahlungs-/Store-Gebühren) in % vom Umsatz", 0.08, PCT, False,
     "Schätzung: Payment-/Store-Gebühren 5–8 % + KI ≈ 0,004 $ je Klassifikation (BP §9). Ergibt J1 ≈ 7 T€ (BP: 8)."),
    ("b2b_month", "Monat im Jahr, in dem B2B-/Pilot-Erlöse zufließen (1–12)", 10, "0", False,
     "Annahme: Pilot-/Lizenzzahlung als Einmalbetrag im Oktober des Jahres (J2 → Monat 22)."),
    ("pack_price", "Ø Preis Creator-Pack (€)", 3.99, '#,##0.00" €"', False, "BP §11.1: 1 Pack à Ø 3,99 €."),
    ("agency_price", "Agentur-Lizenz €/Monat", 49, EUR, False, "BP §11.1: Agentur-Lizenz à 49 €/Monat (Jahresendbestand × 12, wie im Plan)."),
    ("avg_j1", "Ø zahlende Pro im Jahr J1 (Anteil am Jahresendwert)", 0.5, PCT, False, "BP §11.1: ≈ 50 % des Jahresendwerts in J1."),
    ("avg_j2", "Ø zahlende Pro im Jahr J2–J5 (Anteil am Jahresendwert)", 0.75, PCT, False, "BP §11.1: 75 % in J2/J3; für J4/J5 fortgeschrieben."),
    ("buffer_months", "Mindestkasse in Monatskosten (für Folgerunden-Bedarf)", 3, "0", False,
     "Annahme: 3 Monatskosten des Jahres J3 als Sicherheitsbestand."),
    ("defer_share", "Anteil nachgelagerter Vision-Linien an den Vision-Kosten (Story-Engine, Studio/Auto-Edit, Räume/VR-AR)", 0.6, PCT, False,
     "INVESTOR-MEMO §10.1 (500-T€-Fassung): Vision 100 T€ = WortBild 40 + Story-Engine 40 + Highlights/Räume 20 → 60 % nicht WortBild."),
    ("vision_cost_j2", "Vision-Zusatzkosten J2 laut BP §11.5 (€)", 115000, EUR, False, "BP §11.5: Zusatzkosten 30 / 115 / 210 T€."),
    ("vision_cost_j3", "Vision-Zusatzkosten J3 laut BP §11.5 (€)", 210000, EUR, False, "BP §11.5."),
    ("ausbau_start", "Ausbau-Team startet in Monat (nach Gate 2 / Folgerunde)", 19, "0", False,
     "Neu: Einstellungen über das schlanke Pre-Seed-Team hinaus erst nach Gate 2 (Monat 18)."),
]
r += 1
for key, label, val, fmt, keycell, cmt in gen:
    put(A, f"A{r}", label)
    put(A, f"B{r}", val, F_IN, fmt, FILL_KEY if keycell else None, key=key)
    note(A, f"I{r}", cmt)
    r += 1

r += 1
section(A, r, "B. Szenarien (Nutzer, Conversion, Preise, Lizenzen, Kosten J3–J5)", 9)
r += 1
header_row(A, r, ["Annahme", "Einheit", "Konservativ", "Basis", "Optimistisch", "", "", "", "Quelle / Begründung"])
r += 1
scen_rows = [
    ("users_1", "Registrierte Nutzer Ende J1", "Nutzer", (8000, 20000, 40000), NUM, "BP §11.1"),
    ("users_2", "Registrierte Nutzer Ende J2", "Nutzer", (30000, 80000, 180000), NUM, "BP §11.1"),
    ("users_3", "Registrierte Nutzer Ende J3", "Nutzer", (80000, 250000, 500000), NUM, "BP §11.1 (SOM J3: 5.000–15.000 Zahlende, BP §6)"),
    ("ug_4", "Nutzerwachstum J4 (ggü. J3)", "%", (0.5, 0.8, 0.8), PCT, "Neu (Fortschreibung): Wachstum verlangsamt sich ggü. J2→J3 (×2,7 / ×3,1 / ×2,8)."),
    ("ug_5", "Nutzerwachstum J5 (ggü. J4)", "%", (0.35, 0.5, 0.5), PCT, "Neu (Fortschreibung): weitere Verlangsamung; Basis J5 = 675.000 Nutzer ≈ 23–68 % des SAM (1–3 Mio., BP §6)."),
    ("conv_13", "Conversion Free → Pro J1–J3", "%", (0.03, 0.04, 0.05), PCT, "BP §11.1"),
    ("conv_4", "Conversion Free → Pro J4", "%", (0.03, 0.0425, 0.0525), PCT, "Neu: +0,25 Pkt./Jahr in Basis/Optimistisch (Pro+, Packs, Bindung); konservativ konstant."),
    ("conv_5", "Conversion Free → Pro J5", "%", (0.03, 0.045, 0.055), PCT, "Neu: wie oben."),
    ("arpu", "ARPU Pro (netto, Mix Monat/Jahr)", "€/Monat", (8.5, 8.5, 8.5), '#,##0.00" €"', "BP §11.1 (9,99 €/Monat bzw. 79 €/Jahr, netto)"),
    ("pack_share", "Pack-Käufe pro Jahr (Anteil der Nutzer × 1 Pack)", "%", (0.07, 0.10, 0.12), PCT, "BP §11.1"),
    ("ag_1", "Agentur-Lizenzen Ende J1", "Anzahl", (2, 3, 5), NUM, "BP §11.1"),
    ("ag_2", "Agentur-Lizenzen Ende J2", "Anzahl", (8, 15, 30), NUM, "BP §11.1"),
    ("ag_3", "Agentur-Lizenzen Ende J3", "Anzahl", (25, 40, 80), NUM, "BP §11.1"),
    ("ag_4", "Agentur-Lizenzen Ende J4", "Anzahl", (40, 70, 140), NUM, "Neu (Fortschreibung)"),
    ("ag_5", "Agentur-Lizenzen Ende J5", "Anzahl", (60, 100, 200), NUM, "Neu (Fortschreibung)"),
    ("b2b_1", "B2B-Pilot/Lizenz J1", "€", (0, 0, 0), EUR, "BP §11.1"),
    ("b2b_2", "B2B-Pilot/Lizenz J2", "€", (0, 50000, 100000), EUR, "BP §11.1"),
    ("b2b_3", "B2B-Pilot/Lizenz J3", "€", (50000, 150000, 400000), EUR, "BP §11.1"),
    ("b2b_4", "B2B-Lizenz J4", "€", (100000, 300000, 700000), EUR, "Neu (Fortschreibung): 2. Lizenzjahr bzw. 2. Partner"),
    ("b2b_5", "B2B-Lizenz J5", "€", (150000, 500000, 1000000), EUR, "Neu (Fortschreibung)"),
    ("cost_3", "Kosten J3 (nach Folgerunde, wie Plan)", "€", (330000, 710000, 1100000), EUR, "BP §11.3 (Basis 710) und Text darunter (konservativ ≈ 330, optimistisch ≈ 1.100)"),
    ("cg_4", "Kostenwachstum J4", "%", (0.2, 0.5, 0.5), PCT, "Neu (Fortschreibung): Team- und Marketingaufbau"),
    ("cg_5", "Kostenwachstum J5", "%", (0.15, 0.4, 0.4), PCT, "Neu (Fortschreibung)"),
    ("plan_cost_1", "Referenz: Plan-Kosten J1 (BP 500-T€-Fassung)", "€", (120000, 165000, 190000), EUR, "BP §11.3 – nur Vergleich, wird im Modell ersetzt"),
    ("plan_cost_2", "Referenz: Plan-Kosten J2 (BP 500-T€-Fassung)", "€", (230000, 398000, 520000), EUR, "BP §11.3 – Grundlage für den Ausbau ab Monat 19"),
]
for key, label, unit, vals, fmt, src in scen_rows:
    put(A, f"A{r}", label)
    put(A, f"B{r}", unit)
    for i, s in enumerate(SCEN):
        put(A, f"{L(3 + i)}{r}", vals[i], F_IN, fmt, key=f"{key}_{s}")
    note(A, f"I{r}", src)
    r += 1
# abgeleitet: Ausbau-Faktor
put(A, f"A{r}", "Ausbau-Faktor ab Monat 19 (Plan-Kosten J2 Szenario ÷ Basis)")
put(A, f"B{r}", "Faktor")
for i, s in enumerate(SCEN):
    col = L(3 + i)
    put(A, f"{col}{r}", f"={col}{r-1}/$D${r-1}", fmt="0.00", key=f"ausbau_factor_{s}")
note(A, f"I{r}", "Formel: Ausbau-Kosten je Szenario skaliert wie die Plan-Kosten J2 (230 / 398 / 520 T€).")
r += 2

# ---- C. Kostenlinien Pre-Seed
section(A, r, "C. Kostenlinien Pre-Seed (Monat 1–18) – Vorschlag, schlanker Aufbau mit 250.000 €", 9)
r += 1
header_row(A, r, ["Kostenlinie", "Budgetbereich", "€/Monat", "von Monat", "bis Monat", "Einmalbetrag €",
                  "im Monat", "Summe M1–18", "Inhalt / Begründung / ersetzt Planzeile"])
r += 1
cost_lines = [
    ("cl_dev", "Mobile-/Web-Entwickler:in (Vollzeit, ab Monat 3)", "team", 5500, 3, 18, 0, 1,
     "Erste Einstellung (BP §12). 66 T€/Jahr Vollkosten statt 75 T€ J1 / 160 T€ J2 (BP §11.3: 1 statt 2 Entwickler:innen in J2)."),
    ("cl_founder", "Gründergehalt (anteilig)", "team", 2000, 1, 18, 0, 1,
     "BP §11.3: 30 T€ J1 / 48 T€ J2 → 24 T€/Jahr bis Monat 18."),
    ("cl_tools", "Entwicklungs-Tools, Testgeräte (iOS/Android)", "team", 0, 1, 18, 1000, 1, "Einmalig."),
    ("cl_mkt1", "Creator-Programm, Community, Marketing – Phase 1 (TR-Launch)", "gtm", 2000, 1, 6, 0, 1,
     "BP §10 Phase 1; BP §11.3 Marketing 25 T€ J1."),
    ("cl_mkt2", "Creator-Programm, Marketing, Marktplätze – Phase 2 (DE/EN, Partner)", "gtm", 3000, 7, 18, 0, 1,
     "BP §10 Phase 2/3: Botschafter:innen, Pack-Drops, OBS/Streamlabs/Stream-Deck-Einträge."),
    ("cl_pilot", "Plattform- oder Bildungspilot (8 Wochen, Material, Reisen)", "gtm", 0, 1, 18, 2000, 9,
     "Einmalig in Monat 9 (BP §14.1: Plattform-Pilot Q3 2027)."),
    ("cl_wortbild", "Vision-Linie WortBild (Sprachenlernen): Illustration, Sprecher-Audios DE/TR/EN, Didaktik, Kurs-Pilot", "wortbild", 3125, 4, 15, 0, 1,
     "Ersetzt Freelance Design/Sound (BP §11.3: 15 + 25 T€) und den WortBild-Anteil der Vision-Kosten (BP §11.5)."),
    ("cl_legal", "Recht, Marke (DE/EU/TR), Datenschutz/DSFA, Jugendschutz, Steuerberatung", "legal", 1000, 1, 18, 7000, 1,
     "7 T€ einmalig (Marke, Beteiligungsvertrag) + 1 T€/Monat. BP §11.3: 12 + 20 T€."),
]
cl_first = r
for key, label, cat, rate, frm, to, once, once_m, why in cost_lines:
    put(A, f"A{r}", label)
    put(A, f"B{r}", {"team": "Team/Produkt", "gtm": "Markteintritt", "wortbild": "WortBild", "legal": "Recht"}[cat])
    put(A, f"C{r}", rate, F_IN, EUR, key=f"{key}_rate")
    put(A, f"D{r}", frm, F_IN, "0", key=f"{key}_from")
    put(A, f"E{r}", to, F_IN, "0", key=f"{key}_to")
    put(A, f"F{r}", once, F_IN, EUR, key=f"{key}_once")
    put(A, f"G{r}", once_m, F_IN, "0", key=f"{key}_oncem")
    put(A, f"H{r}", f"=C{r}*MAX(0,MIN(E{r},18)-D{r}+1)+IF(G{r}<=18,F{r},0)", fmt=EUR, key=f"{key}_sum18")
    note(A, f"I{r}", why)
    A[f"B{r}"].font = Font(name=FONT, color="000000")
    r += 1
cl_last = r - 1
put(A, f"A{r}", "Summe verplante Pre-Seed-Kosten M1–18", bold=True)
put(A, f"H{r}", f"=SUM(H{cl_first}:H{cl_last})", fmt=EUR, bold=True, key="cl_total18")
r += 1
put(A, f"A{r}", "Laufende Kosten nach Monat 18 im Fall „ohne Umsatz“ (schlank weiter: Entwicklung, Gründer, Marketing Ph. 2, Recht)")
put(A, f"C{r}", f"={R['cl_dev_rate'][1]}+{R['cl_founder_rate'][1]}+{R['cl_mkt2_rate'][1]}+{R['cl_legal_rate'][1]}",
    fmt=EUR, key="lean_rate")
note(A, f"I{r}", "Annahme für die Runway ohne Umsatz: keine Neueinstellungen, Pre-Seed-Team läuft weiter.")
r += 2

# ---- D. Ausbau ab Monat 19 (Basis = BP 11.3 J2 / 12)
section(A, r, "D. Ausbau-Team ab Monat 19 (nach Gate 2) – Basis = BP §11.3 J2 ÷ 12, andere Szenarien × Ausbau-Faktor", 9)
r += 1
header_row(A, r, ["Kostenlinie (BP §11.3)", "Kategorie", "Plan J2 €/Jahr", "€/Monat ab M19", "", "", "", "", "Hinweis"])
r += 1
ausbau = [
    ("ab_dev", "Personal: Mobile-/Web-Entwicklung (2 Personen)", "Mobile/Web", 160000),
    ("ab_backend", "Personal: Backend/ML, Community/Support", "Backend/Community", 60000),
    ("ab_founder", "Gründer (Gehalt)", "Gründer", 48000),
    ("ab_free", "Freelance Design, Sound, Illustration (Packs, WortBild)", "Freelance/WortBild", 25000),
    ("ab_mkt", "Marketing, Creator-Partnerprogramm, Events", "Marketing", 60000),
    ("ab_legal", "Recht, Marke, Steuer, Verwaltung", "Recht", 20000),
]
ab_first = r
for key, label, cat, val in ausbau:
    put(A, f"A{r}", label)
    put(A, f"B{r}", cat)
    put(A, f"C{r}", val, F_IN, EUR, key=f"{key}_plan")
    put(A, f"D{r}", f"=C{r}/12", fmt=EUR, key=f"{key}_rate")
    r += 1
put(A, f"A{r}", "Summe Ausbau Basis €/Monat (ohne Infrastruktur)", bold=True)
put(A, f"C{r}", f"=SUM(C{ab_first}:C{r-1})", fmt=EUR, bold=True)
put(A, f"D{r}", f"=SUM(D{ab_first}:D{r-1})", fmt=EUR, bold=True, key="ab_total_rate")
note(A, f"I{r}", "Infrastruktur läuft separat (fix + % vom Umsatz). Plan J2 gesamt 398 T€ = 373 T€ + Infrastruktur 25 T€.")
r += 2

# ---- E. Plan-Referenz Umsatz und Ergebnis (BP 11.2 / 11.4)
section(A, r, "E. Referenz: Plan-Werte BP §11.2 / §11.4 (T€, 500-T€-Fassung) – nur zum Abgleich", 9)
r += 1
header_row(A, r, ["Zeile", "Szenario", "J1", "J2", "J3", "", "", "", "Hinweis"])
r += 1
plan_ref = {
    "kons": {"rev": (15, 81, 268), "ag": (1, 4, 12), "res": (-105, -149, -62)},
    "basis": {"rev": (51, 336, 1039), "ag": (2, 9, 24), "res": (-114, -62, 329)},
    "opt": {"rev": (124, 893, 2601), "ag": (3, 18, 48), "res": (-66, 373, 1501)},
}
for s in SCEN:
    for k, lab in (("rev", "Umsatz gesamt (Plan)"), ("ag", "davon Agentur-Lizenz (Plan)"), ("res", "Ergebnis (Plan)")):
        put(A, f"A{r}", lab)
        put(A, f"B{r}", SCEN_DE[s])
        for j in range(3):
            put(A, f"{L(3 + j)}{r}", plan_ref[s][k][j], F_IN, NUM, key=f"plan_{k}_{s}_{j+1}")
        if k == "ag" and s == "kons":
            note(A, f"I{r}", "Inkonsistent: 8 × 49 € × 12 = 4,7 T€ und 25 × 49 € × 12 = 14,7 T€ – Modell rechnet mit der Formel (korrigiert).")
        r += 1
r += 1

# ---- F. Änderungen ggü. Plan
section(A, r, "F. Geänderte Annahmen gegenüber BUSINESSPLAN.md (500-T€-Fassung) – dokumentiert", 9)
r += 1
changes = [
    "1. Rundengröße 250.000 € statt 500.000 €; Mittelverwendung über 18 statt 24 Monate (Blatt Mittelverwendung).",
    "2. Kosten J1–J2 aus dem Monatsplan statt BP §11.3: Pre-Seed-Team (1 Entwickler:in ab M3, Gründergehalt 2 T€/Monat) bis M18; "
    "Ausbau-Team gemäß Plan J2 erst ab Monat 19 nach Gate 2 (Backend/ML, Community, 2. Entwickler:in).",
    "3. Freelance Design/Sound (BP §11.3) geht in der Linie WortBild auf (37,5 T€, M4–M15); Story-Engine, Auto-Edit/Studio und VR/AR "
    "wandern in die Folgerunde (60 % der Vision-Kosten BP §11.5 als Bedarf der Folgerunde).",
    "4. Recht/Marke/Datenschutz 25 T€ in 18 Monaten (Plan: 12 T€ J1 + 20 T€ J2).",
    "5. Infrastruktur als 250 €/Monat fix + 8 % vom Umsatz statt fester Jahreswerte (J1 ≈ 7 T€ statt 8 T€).",
    "6. Umsatz J1–J3 unverändert aus BP §11.1 berechnet; Korrektur Agentur-Zeile konservativ (J2 4,7 statt 4 T€, J3 14,7 statt 12 T€); "
    "übrige Abweichungen ≤ 1 T€ sind Rundung. Vision-Erlöse (BP §11.5) bleiben – wie im Plan – außerhalb des Ergebnisses.",
    "7. J4–J5 neu fortgeschrieben (Nutzerwachstum, Conversion, Agenturen, B2B, Kostenwachstum – Abschnitt B).",
    "8. Monatlicher Umsatzverlauf: J1 linear ab null (Summe = Jahreswert), J2 linear ab dem Dezember-Niveau J1 (Summe = Jahreswert); "
    "B2B als Einmalzahlung im Monat laut Abschnitt A.",
]
for t in changes:
    put(A, f"A{r}", t)
    r += 1

A.freeze_panes = "A5"

# ------------------------------------------------------------------ Blatt: Liquidität
Q = wb.create_sheet("Liquidität 24M")
widths(Q, {"A": 50, "B": 18, "C": 16, "D": 16, "E": 16, "F": 16, "G": 40})
for m in range(1, 25):
    Q.column_dimensions[L(3 + m)].width = 13
for col in ("AB", "AC", "AD"):
    Q.column_dimensions[col].width = 13
MC = {m: L(3 + m) for m in range(1, 25)}  # Monat -> Spalte (D..AA)
put(Q, "A1", "Liquidität monatlich, 24 Monate – Startkasse 250.000 € (Closing Jan 2027)", F_TITLE)
note(Q, "A2", DISCLAIMER)

# KPI-Block oben (Zeilen 4–24), Monatsplan ab Zeile 27
KPI_TOP = 4
MROW = 27  # Kopfzeile Monate
put(Q, f"A{MROW}", "Monat", bold=True)
put(Q, f"A{MROW+1}", "Kalendermonat", bold=True)
for m in range(1, 25):
    put(Q, f"{MC[m]}{MROW}", m, F_BOLD, "0")
    yr = 2027 + (m - 1) // 12
    mo = (m - 1) % 12 + 1
    put(Q, f"{MC[m]}{MROW+1}", f"{mo:02d}/{yr}", align="center")
put(Q, f"AB{MROW}", "Summe J1", F_BOLD)
put(Q, f"AC{MROW}", "Summe J2", F_BOLD)
put(Q, f"AD{MROW}", "Summe M1–18", F_BOLD)
mrow_ref = lambda col: f"{col}${MROW}"

def month_sums(row, fmt=EUR):
    put(Q, f"AB{row}", f"=SUM(D{row}:O{row})", fmt=fmt)
    put(Q, f"AC{row}", f"=SUM(P{row}:AA{row})", fmt=fmt)
    put(Q, f"AD{row}", f"=SUM(D{row}:U{row})", fmt=fmt)

r = MROW + 3
section(Q, r, "Kosten Pre-Seed (Monat 1–18, in allen Szenarien gleich) – aus Annahmen C", 30)
r += 1
rows_cl = {}
for key, label, *_ in cost_lines:
    put(Q, f"A{r}", label)
    put(Q, f"B{r}", "€")
    for m in range(1, 25):
        c = MC[m]
        f = (f"=IF(AND({c}${MROW}>={ref(key+'_from')},{c}${MROW}<={ref(key+'_to')}),{ref(key+'_rate')},0)"
             f"+IF({c}${MROW}={ref(key+'_oncem')},{ref(key+'_once')},0)")
        put(Q, f"{c}{r}", f, F_LINK, EUR)
    month_sums(r)
    rows_cl[key] = r
    r += 1
ROW_PRESEED = r
put(Q, f"A{r}", "Summe Pre-Seed-Kosten", bold=True)
for m in range(1, 25):
    c = MC[m]
    put(Q, f"{c}{r}", f"=SUM({c}{r-len(cost_lines)}:{c}{r-1})", fmt=EUR, bold=True, fill=FILL_TOT)
month_sums(r)
r += 2

section(Q, r, "Ausbau-Team ab Monat 19 (Basis; Szenarien × Ausbau-Faktor) – aus Annahmen D", 30)
r += 1
rows_ab = {}
for key, label, cat, val in ausbau:
    put(Q, f"A{r}", label)
    put(Q, f"B{r}", cat)
    for m in range(1, 25):
        c = MC[m]
        put(Q, f"{c}{r}", f"=IF({c}${MROW}>={ref('ausbau_start')},{ref(key+'_rate')},0)", F_LINK, EUR)
    month_sums(r)
    rows_ab[key] = r
    r += 1
ROW_AUSBAU = r
put(Q, f"A{r}", "Summe Ausbau (Basis)", bold=True)
for m in range(1, 25):
    c = MC[m]
    put(Q, f"{c}{r}", f"=SUM({c}{r-len(ausbau)}:{c}{r-1})", fmt=EUR, bold=True, fill=FILL_TOT)
month_sums(r)
r += 2

ROW_LEAN = r
put(Q, f"A{r}", "Weiterbetrieb schlank ab Monat 19 (nur Fall „ohne Umsatz“)")
for m in range(1, 25):
    c = MC[m]
    put(Q, f"{c}{r}", f"=IF({c}${MROW}>={ref('ausbau_start')},{ref('lean_rate')},0)", F_LINK, EUR)
month_sums(r)
r += 2

# Szenario-Blöcke
SR = {}  # s -> dict row names
for s in SCEN:
    d = {}
    section(Q, r, f"Szenario {SCEN_DE[s]}", 30)
    r += 1
    # Hilfswerte (Spalte C)
    helpers = [
        ("rec1", "Wiederkehrender Umsatz J1 (Pro + Packs + Agentur), Jahr", f"='GuV 5 Jahre'!$C${{rec_{s}}}"),
        ("rec2", "Wiederkehrender Umsatz J2, Jahr", f"='GuV 5 Jahre'!$D${{rec_{s}}}"),
        ("s0", "Startniveau J2 = Dezember-Umsatz J1 (€/Monat)", None),
        ("g", "Monatlicher Zuwachs J2 (€/Monat)", None),
        ("b2b1", "B2B J1 (Einmalzahlung)", None),
        ("b2b2", "B2B J2 (Einmalzahlung)", None),
    ]
    for hk, lab, f in helpers:
        put(Q, f"A{r}", lab)
        d[hk] = r
        r += 1
    d["rev_rec"] = r; r += 1
    d["rev_b2b"] = r; r += 1
    d["rev"] = r; r += 1
    d["c_pre"] = r; r += 1
    d["c_ab"] = r; r += 1
    d["c_inf"] = r; r += 1
    d["cost"] = r; r += 1
    d["net"] = r; r += 1
    d["cash"] = r; r += 1
    d["f_neg"] = r; r += 1
    d["f_be"] = r; r += 1
    d["f_sus"] = r; r += 1
    d["f_buf"] = r; r += 1
    r += 1
    SR[s] = d
# Ohne Umsatz
section(Q, r, "Fall „ohne jeden Umsatz“ (Runway-Test)", 30)
r += 1
NR = {"cost": r, "cash": r + 1, "f_neg": r + 2}
r += 4
LAST_Q = r


# ------------------------------------------------------------------ Blatt: GuV
G = wb.create_sheet("GuV 5 Jahre", 1)
widths(G, {"A": 52, "B": 12, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 60})
put(G, "A1", "GuV 5 Jahre (EBITDA-Sicht) – drei Szenarien", F_TITLE)
note(G, "A2", DISCLAIMER + " J1–J3 aus BP §11.1, Kosten J1–J2 aus dem Monatsplan (250-T€-Runde), J4–J5 fortgeschrieben.")
r = 4
GR = {}
for s in SCEN:
    d = {}
    section(G, r, f"Szenario {SCEN_DE[s]}", 8)
    r += 1
    header_row(G, r, ["Position", "Einheit"] + YEARS + ["Herleitung"])
    r += 1
    def yc(j):
        return L(3 + j)  # j=0..4 -> C..G
    # Nutzer
    d["users"] = r
    put(G, f"A{r}", "Registrierte Nutzer (Jahresende)")
    put(G, f"B{r}", "Nutzer")
    for j in range(5):
        c = yc(j)
        if j < 3:
            f = f"={ref(f'users_{j+1}_{s}')}"
        else:
            f = f"={yc(j-1)}{r}*(1+{ref(f'ug_{j+1}_{s}')})"
        put(G, f"{c}{r}", f, fmt=NUM)
    note(G, f"H{r}", "J1–J3 BP §11.1; J4–J5 = Vorjahr × (1 + Wachstum)")
    r += 1
    d["conv"] = r
    put(G, f"A{r}", "Conversion Free → Pro")
    put(G, f"B{r}", "%")
    for j in range(5):
        k = "conv_13" if j < 3 else f"conv_{j+1}"
        put(G, f"{yc(j)}{r}", f"={ref(f'{k}_{s}')}", fmt=PCT)
    r += 1
    d["pro_end"] = r
    put(G, f"A{r}", "Zahlende Pro-Nutzer (Jahresende)")
    put(G, f"B{r}", "Abos")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['users']}*{c}{d['conv']}", fmt=NUM)
    r += 1
    d["pro_avg"] = r
    put(G, f"A{r}", "Zahlende Pro-Nutzer (Jahresdurchschnitt)")
    put(G, f"B{r}", "Abos")
    for j in range(5):
        c = yc(j)
        fac = ref("avg_j1") if j == 0 else ref("avg_j2")
        put(G, f"{c}{r}", f"={c}{d['pro_end']}*{fac}", fmt=NUM)
    note(G, f"H{r}", "BP §11.1: 50 % des Jahresendwerts in J1, 75 % ab J2")
    r += 1
    d["r_pro"] = r
    put(G, f"A{r}", "Umsatz Pro-Abo")
    put(G, f"B{r}", "€")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['pro_avg']}*{ref(f'arpu_{s}')}*12", fmt=EUR)
    r += 1
    d["r_pack"] = r
    put(G, f"A{r}", "Umsatz Creator-Packs")
    put(G, f"B{r}", "€")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['users']}*{ref(f'pack_share_{s}')}*{ref('pack_price')}", fmt=EUR)
    r += 1
    d["r_ag"] = r
    put(G, f"A{r}", "Umsatz Agentur-Lizenz")
    put(G, f"B{r}", "€")
    for j in range(5):
        put(G, f"{yc(j)}{r}", f"={ref(f'ag_{j+1}_{s}')}*{ref('agency_price')}*12", fmt=EUR)
    note(G, f"H{r}", "Jahresendbestand × 49 € × 12 (Methode des Plans)")
    r += 1
    d["rec"] = r
    put(G, f"A{r}", "Zwischensumme wiederkehrender Umsatz")
    put(G, f"B{r}", "€")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['r_pro']}+{c}{d['r_pack']}+{c}{d['r_ag']}", fmt=EUR)
    r += 1
    d["r_b2b"] = r
    put(G, f"A{r}", "Umsatz B2B / Pilot")
    put(G, f"B{r}", "€")
    for j in range(5):
        put(G, f"{yc(j)}{r}", f"={ref(f'b2b_{j+1}_{s}')}", fmt=EUR)
    r += 1
    d["rev"] = r
    put(G, f"A{r}", "Umsatz gesamt", bold=True)
    put(G, f"B{r}", "€")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['rec']}+{c}{d['r_b2b']}", fmt=EUR, bold=True, fill=FILL_TOT, key=f"guv_rev_{s}_{j+1}")
    r += 1
    d["cost"] = r
    put(G, f"A{r}", "Kosten gesamt (operativ)", bold=True)
    put(G, f"B{r}", "€")
    q = SR[s]
    put(G, f"C{r}", f"='Liquidität 24M'!$AB${q['cost']}", fmt=EUR, bold=True, key=f"guv_cost_{s}_1")
    put(G, f"D{r}", f"='Liquidität 24M'!$AC${q['cost']}", fmt=EUR, bold=True, key=f"guv_cost_{s}_2")
    put(G, f"E{r}", f"={ref(f'cost_3_{s}')}", fmt=EUR, bold=True, key=f"guv_cost_{s}_3")
    put(G, f"F{r}", f"=E{r}*(1+{ref(f'cg_4_{s}')})", fmt=EUR, bold=True, key=f"guv_cost_{s}_4")
    put(G, f"G{r}", f"=F{r}*(1+{ref(f'cg_5_{s}')})", fmt=EUR, bold=True, key=f"guv_cost_{s}_5")
    note(G, f"H{r}", "J1–J2 = Summe Monatsplan (Liquidität 24M); J3 = BP §11.3; J4–J5 = Vorjahr × (1 + Kostenwachstum)")
    r += 1
    d["ebitda"] = r
    put(G, f"A{r}", "EBITDA (Ergebnis)", bold=True)
    put(G, f"B{r}", "€")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['rev']}-{c}{d['cost']}", fmt=EUR, bold=True, fill=FILL_KEY, key=f"guv_ebitda_{s}_{j+1}")
    r += 1
    d["margin"] = r
    put(G, f"A{r}", "EBITDA-Marge")
    put(G, f"B{r}", "%")
    for j in range(5):
        c = yc(j)
        put(G, f"{c}{r}", f"=IF({c}{d['rev']}=0,0,{c}{d['ebitda']}/{c}{d['rev']})", fmt=PCT)
    r += 1
    d["cum"] = r
    put(G, f"A{r}", "Kumuliertes Ergebnis")
    put(G, f"B{r}", "€")
    for j in range(5):
        c = yc(j)
        f = f"={c}{d['ebitda']}" if j == 0 else f"={yc(j-1)}{r}+{c}{d['ebitda']}"
        put(G, f"{c}{r}", f, fmt=EUR, key=f"guv_cum_{s}_{j+1}")
    r += 1
    d["cash"] = r
    put(G, f"A{r}", "Kasse Jahresende (250 T€ Start, ohne Folgerunde)")
    put(G, f"B{r}", "€")
    put(G, f"C{r}", f"='Liquidität 24M'!$O${q['cash']}", fmt=EUR, key=f"guv_cash_{s}_1")
    put(G, f"D{r}", f"='Liquidität 24M'!$AA${q['cash']}", fmt=EUR, key=f"guv_cash_{s}_2")
    for j in (2, 3, 4):
        c = yc(j)
        put(G, f"{c}{r}", f"={yc(j-1)}{r}+{c}{d['ebitda']}", fmt=EUR, key=f"guv_cash_{s}_{j+1}")
    note(G, f"H{r}", "J1/J2 aus dem Monatsplan (Monat 12/24); ab J3 Vorjahr + EBITDA (vereinfacht, ohne Steuern/Working Capital)")
    r += 1
    # Plan-Abgleich
    d["plan_rev"] = r
    put(G, f"A{r}", "Referenz: Umsatz laut Plan BP §11.2 (T€)", F_NOTE)
    for j in range(3):
        put(G, f"{yc(j)}{r}", f"={ref(f'plan_rev_{s}_{j+1}')}", F_LINK, NUM)
    r += 1
    put(G, f"A{r}", "Abweichung Modell − Plan beim Umsatz (T€)", F_NOTE)
    for j in range(3):
        c = yc(j)
        put(G, f"{c}{r}", f"={c}{d['rev']}/1000-{c}{d['plan_rev']}", fmt='0.0;(0.0);"–"')
    note(G, f"H{r}", "Rundung; konservativ J2/J3: Korrektur der Agentur-Zeile (siehe Annahmen F)")
    r += 1
    put(G, f"A{r}", "Referenz: Ergebnis laut Plan BP §11.4 (T€, 500-T€-Kostenbasis)", F_NOTE)
    for j in range(3):
        put(G, f"{yc(j)}{r}", f"={ref(f'plan_res_{s}_{j+1}')}", F_LINK, NUM)
    r += 2
    GR[s] = d

# WortBild-Memo
section(G, r, "Memo: Vision-Linie WortBild – Erlöse laut BP §11.5 (Basis, NICHT im Ergebnis enthalten)", 8)
r += 1
header_row(G, r, ["Position", "Einheit"] + YEARS[:3] + ["", "", "Quelle"])
r += 1
wb_rows = [("WortBild Familie (Ø 0 / 500 / 2.000 Abos × 4,25 € × 12)", (0, 26000, 102000)),
           ("Schullizenzen (5 / 40 / 150 Schulen × Ø 500 €)", (2500, 20000, 75000)),
           ("Kurslizenzen (0 / 200 / 800 Lehrkräfte × 49 €)", (0, 9800, 39200))]
wb_first = r
for lab, vals in wb_rows:
    put(G, f"A{r}", lab)
    put(G, f"B{r}", "€")
    for j in range(3):
        put(G, f"{L(3+j)}{r}", vals[j], F_IN, EUR)
    note(G, f"H{r}", "BP §11.5 (Schätzung)")
    r += 1
put(G, f"A{r}", "Summe WortBild-Erlöse (Option, Upside)", bold=True)
for j in range(3):
    c = L(3 + j)
    put(G, f"{c}{r}", f"=SUM({c}{wb_first}:{c}{r-1})", fmt=EUR, bold=True, key=f"wortbild_rev_{j+1}")
r += 2
G.freeze_panes = "C4"

# --- jetzt die Liquiditäts-Formeln füllen (brauchen GuV-Zeilen)
for s in SCEN:
    d = SR[s]
    g = GR[s]
    put(Q, f"C{d['rec1']}", f"='GuV 5 Jahre'!$C${g['rec']}", fmt=EUR)
    put(Q, f"C{d['rec2']}", f"='GuV 5 Jahre'!$D${g['rec']}", fmt=EUR)
    put(Q, f"C{d['s0']}", f"=C{d['rec1']}*11.5/72", fmt=EUR)
    put(Q, f"C{d['g']}", f"=(C{d['rec2']}-12*C{d['s0']})/78", fmt=EUR)
    put(Q, f"C{d['b2b1']}", f"='GuV 5 Jahre'!$C${g['r_b2b']}", fmt=EUR)
    put(Q, f"C{d['b2b2']}", f"='GuV 5 Jahre'!$D${g['r_b2b']}", fmt=EUR)
    note(Q, f"D{d['s0']}", "J1: Monat m = Jahreswert × (m − 0,5) ÷ 72 (linear ab null, Summe = Jahreswert)")
    note(Q, f"D{d['g']}", "J2: Monat m = s + g × (m − 12); g so gewählt, dass die Summe J2 dem Jahreswert entspricht")
    labels = {
        "rev_rec": "Umsatz wiederkehrend (Pro, Packs, Agentur)", "rev_b2b": "Umsatz B2B/Pilot (Einmalzahlung)",
        "rev": "Umsatz gesamt", "c_pre": "Kosten Pre-Seed-Team und -Budget", "c_ab": "Kosten Ausbau-Team (ab M19)",
        "c_inf": "Infrastruktur (fix + % vom Umsatz)", "cost": "Kosten gesamt", "net": "Netto-Cashflow",
        "cash": "Kasse Monatsende", "f_neg": "Hilfszeile: Kasse < 0 (1 = ja)", "f_be": "Hilfszeile: Umsatz ≥ Kosten",
        "f_sus": "Hilfszeile: Umsatz ≥ Kosten ab hier bis M24", "f_buf": "Hilfszeile: Kasse < Mindestkasse (3 Monatskosten)",
    }
    for k, lab in labels.items():
        bold = k in ("rev", "cost", "cash")
        put(Q, f"A{d[k]}", lab, bold=bold, font=F_NOTE if k.startswith("f_") else None)
    for m in range(1, 25):
        c = MC[m]
        mm = f"{c}${MROW}"
        put(Q, f"{c}{d['rev_rec']}", f"=IF({mm}<=12,$C${d['rec1']}*({mm}-0.5)/72,$C${d['s0']}+$C${d['g']}*({mm}-12))", fmt=EUR)
        put(Q, f"{c}{d['rev_b2b']}", f"=IF({mm}={ref('b2b_month')},$C${d['b2b1']},0)+IF({mm}=12+{ref('b2b_month')},$C${d['b2b2']},0)", fmt=EUR)
        put(Q, f"{c}{d['rev']}", f"={c}{d['rev_rec']}+{c}{d['rev_b2b']}", fmt=EUR, bold=True)
        put(Q, f"{c}{d['c_pre']}", f"={c}${ROW_PRESEED}", fmt=EUR)
        put(Q, f"{c}{d['c_ab']}", f"={c}${ROW_AUSBAU}*{ref(f'ausbau_factor_{s}')}", fmt=EUR)
        put(Q, f"{c}{d['c_inf']}", f"={ref('infra_fix')}+{ref('infra_var')}*{c}{d['rev']}", fmt=EUR)
        put(Q, f"{c}{d['cost']}", f"={c}{d['c_pre']}+{c}{d['c_ab']}+{c}{d['c_inf']}", fmt=EUR, bold=True)
        put(Q, f"{c}{d['net']}", f"={c}{d['rev']}-{c}{d['cost']}", fmt=EUR)
        prev = ref("start_cash") if m == 1 else f"{MC[m-1]}{d['cash']}"
        put(Q, f"{c}{d['cash']}", f"={prev}+{c}{d['net']}", fmt=EUR, bold=True, fill=FILL_TOT)
        put(Q, f"{c}{d['f_neg']}", f"=IF({c}{d['cash']}<0,1,0)", F_NOTE, "0")
        put(Q, f"{c}{d['f_be']}", f"=IF({c}{d['rev']}>={c}{d['cost']},1,0)", F_NOTE, "0")
        nxt = "1" if m == 24 else f"{MC[m+1]}{d['f_sus']}"
        put(Q, f"{c}{d['f_sus']}", f"=IF(AND({c}{d['f_be']}=1,{nxt}=1),1,0)", F_NOTE, "0")
        put(Q, f"{c}{d['f_buf']}", f"=IF({c}{d['cash']}<{ref('buffer_months')}*{c}{d['cost']},1,0)", F_NOTE, "0")
    for k in ("rev_rec", "rev_b2b", "rev", "c_pre", "c_ab", "c_inf", "cost", "net"):
        month_sums(d[k])

# ohne Umsatz
put(Q, f"A{NR['cost']}", "Kosten (Pre-Seed bis M18, danach schlank, Infrastruktur fix)", bold=True)
put(Q, f"A{NR['cash']}", "Kasse Monatsende (ohne Umsatz)", bold=True)
put(Q, f"A{NR['f_neg']}", "Hilfszeile: Kasse < 0", F_NOTE)
for m in range(1, 25):
    c = MC[m]
    put(Q, f"{c}{NR['cost']}", f"={c}${ROW_PRESEED}+{c}${ROW_LEAN}+{ref('infra_fix')}", fmt=EUR)
    prev = ref("start_cash") if m == 1 else f"{MC[m-1]}{NR['cash']}"
    put(Q, f"{c}{NR['cash']}", f"={prev}-{c}{NR['cost']}", fmt=EUR, bold=True, fill=FILL_TOT)
    put(Q, f"{c}{NR['f_neg']}", f"=IF({c}{NR['cash']}<0,1,0)", F_NOTE, "0")
month_sums(NR["cost"])

# KPI-Block
r = KPI_TOP
section(Q, r, "Kennzahlen (Ziele/Schätzungen) – Runway, Kasse, Break-even, Folgerunde", 7)
r += 1
header_row(Q, r, ["Kennzahl", "Einheit", "Konservativ", "Basis", "Optimistisch", "Ohne Umsatz", "Formel / Lesart"])
r += 1
def rng(row):
    return f"$D${row}:$AA${row}"
cols4 = {"kons": "C", "basis": "D", "opt": "E", "nr": "F"}
kpis = [
    ("cash_m6", "Kasse Ende Monat 6", "€", lambda d: f"=INDEX({rng(d['cash'])},6)", EUR, "INDEX in die Kassenzeile"),
    ("cash_m12", "Kasse Ende Monat 12", "€", lambda d: f"=INDEX({rng(d['cash'])},12)", EUR, ""),
    ("cash_m18", "Kasse Ende Monat 18", "€", lambda d: f"=INDEX({rng(d['cash'])},18)", EUR, ""),
    ("cash_m24", "Kasse Ende Monat 24", "€", lambda d: f"=INDEX({rng(d['cash'])},24)", EUR, ""),
    ("cash_min", "Niedrigster Kassenstand M1–24", "€", lambda d: f"=MIN({rng(d['cash'])})", EUR, ""),
    ("cash_min_m", "Monat des niedrigsten Kassenstands", "Monat", lambda d: f"=MATCH(MIN({rng(d['cash'])}),{rng(d['cash'])},0)", "0", ""),
    ("cashout_m", "Cash-out-Monat (erster Monat mit Kasse < 0)", "Monat", lambda d: f'=IFERROR(MATCH(1,{rng(d["f_neg"])},0),"keiner bis M24")', "0", "Text = Kasse bleibt bis Monat 24 positiv"),
    ("runway", "Runway in Monaten (bis Kasse < 0)", "Monate", lambda d: f'=IFERROR(MATCH(1,{rng(d["f_neg"])},0)-1,"> 24")', "0", "Monate mit positiver Kasse"),
]
KR = {}
for key, lab, unit, fn, fmt, hint in kpis:
    put(Q, f"A{r}", lab)
    put(Q, f"B{r}", unit)
    for s in SCEN:
        put(Q, f"{cols4[s]}{r}", fn(SR[s]), fmt=fmt, key=f"q_{key}_{s}")
    put(Q, f"F{r}", fn(NR), fmt=fmt, key=f"q_{key}_nr")
    note(Q, f"G{r}", hint)
    KR[key] = r
    r += 1
# nur Szenarien
put(Q, f"A{r}", "Nachhaltiger Break-even (ab diesem Monat Umsatz ≥ Kosten bis M24)")
put(Q, f"B{r}", "Monat")
for s in SCEN:
    put(Q, f"{cols4[s]}{r}", f'=IFERROR(MATCH(1,{rng(SR[s]["f_sus"])},0),"nach M24")', fmt="0", key=f"q_be_{s}")
note(Q, f"G{r}", "operativer Monats-Break-even inkl. Ausbau-Team ab M19")
KR["be"] = r
r += 1
put(Q, f"A{r}", "Ø Kosten pro Monat M1–18 (Burn ohne Umsatz)")
put(Q, f"B{r}", "€/Monat")
put(Q, f"F{r}", f"=SUM($D${NR['cost']}:$U${NR['cost']})/18", fmt=EUR, key="q_burn18")
r += 1
put(Q, f"A{r}", "Kasse fällt unter Mindestkasse (3 laufende Monatskosten) in Monat")
put(Q, f"B{r}", "Monat")
for s in SCEN:
    put(Q, f"{cols4[s]}{r}", f'=IFERROR(MATCH(1,{rng(SR[s]["f_buf"])},0),"nie bis M24")', fmt="0", key=f"q_bufm_{s}")
KR["bufm"] = r
r += 1
put(Q, f"A{r}", "Kasse Ende J3 ohne Folgerunde (Kasse M24 + EBITDA J3)")
put(Q, f"B{r}", "€")
for s in SCEN:
    put(Q, f"{cols4[s]}{r}", f"={cols4[s]}{KR['cash_m24']}+'GuV 5 Jahre'!$E${GR[s]['ebitda']}", fmt=EUR, key=f"q_cash_j3_{s}")
KR["cash_j3"] = r
r += 1
put(Q, f"A{r}", "Mindestkasse = 3 Monatskosten J3")
put(Q, f"B{r}", "€")
for s in SCEN:
    put(Q, f"{cols4[s]}{r}", f"={ref('buffer_months')}*'GuV 5 Jahre'!$E${GR[s]['cost']}/12", fmt=EUR, key=f"q_buffer_{s}")
KR["buffer"] = r
r += 1
put(Q, f"A{r}", "Liquiditätslücke bis Ende J3 = MAX(0; Mindestkasse − niedrigster Stand)")
put(Q, f"B{r}", "€")
for s in SCEN:
    c = cols4[s]
    put(Q, f"{c}{r}", f"=MAX(0,{c}{KR['buffer']}-MIN({c}{KR['cash_min']},{c}{KR['cash_j3']}))", fmt=EUR, key=f"q_gap_{s}")
KR["gap"] = r
r += 1
put(Q, f"A{r}", "Nachgelagerte Linien (Story-Engine, Studio/Auto-Edit, VR/AR) ab M19 bis Ende J3")
put(Q, f"B{r}", "€")
for s in SCEN:
    put(Q, f"{cols4[s]}{r}", f"={ref('defer_share')}*({ref('vision_cost_j2')}/2+{ref('vision_cost_j3')})", fmt=EUR, key=f"q_defer_{s}")
note(Q, f"G{r}", "60 % der Vision-Kosten BP §11.5 (½ J2 + J3)")
KR["defer"] = r
r += 1
put(Q, f"A{r}", "Folgerunde (Seed) – Mindestbedarf", bold=True)
put(Q, f"B{r}", "€")
for s in SCEN:
    c = cols4[s]
    put(Q, f"{c}{r}", f"={c}{KR['gap']}+{c}{KR['defer']}", fmt=EUR, bold=True, fill=FILL_KEY, key=f"q_followon_{s}")
KR["followon"] = r
r += 1
put(Q, f"A{r}", "Folgerunde spätestens abschließen bis Monat (Gate 2 = Monat 18)")
put(Q, f"B{r}", "Monat")
for s in SCEN:
    c = cols4[s]
    put(Q, f"{c}{r}", f"=IF(ISNUMBER({c}{KR['bufm']}),MIN(18,{c}{KR['bufm']}),18)", fmt="0", key=f"q_followon_m_{s}")
note(Q, f"G{r}", "frühestens wenn Gate 2 erfüllt; Gespräche ab Monat 12–15 beginnen")
KR["followon_m"] = r
r += 1
put(Q, f"A{r}", "Seed-Richtwert für die Planung (deckt das konservative Szenario, auf 50 T€ gerundet)", bold=True)
put(Q, f"B{r}", "€")
put(Q, f"D{r}", f"=CEILING(MAX(C{KR['followon']}:E{KR['followon']}),50000)", fmt=EUR, bold=True, fill=FILL_KEY, key="q_seed_guide")
r += 1
Q.freeze_panes = Q[f"D{MROW+2}"]

# ------------------------------------------------------------------ Blatt: Mittelverwendung
M = wb.create_sheet("Mittelverwendung", 2)
widths(M, {"A": 54, "B": 11, "C": 15, "D": 17, "E": 15, "F": 15, "G": 70})
put(M, "A1", "Mittelverwendung 250.000 € über 18 Monate (Vorschlag)", F_TITLE)
note(M, "A2", "Vorschlag – mit Steuer-/Rechtsberater prüfen. Story-Engine, Auto-Edit/Studio und VR/AR werden aus der Folgerunde finanziert.")
r = 4
header_row(M, r, ["Bereich", "Anteil", "Betrag", "Verplant M1–18 (Monatsplan)", "Differenz", "€/Monat Ø", "Inhalt"])
r += 1
uof = [
    ("uof_team", "Team/Produkt: Mobile-/Web-Entwickler:in, Gründergehalt anteilig, Testgeräte", 0.50, ["cl_dev", "cl_founder", "cl_tools"],
     "Entwickler:in ab Monat 3 (16 × 5.500 € = 88 T€), Gründer 18 × 2.000 € = 36 T€, Testgeräte 1 T€; Ziel: Pro-Abo, Mobile-App/PWA, Stabilität v2.1"),
    ("uof_gtm", "Markteintritt: Creator-Programm, Community, Plattform- und Bildungspiloten", 0.20, ["cl_mkt1", "cl_mkt2", "cl_pilot"],
     "TR-Launch, 10–20 Botschafter:innen, Pack-Drops, Marktplatz-Einträge, DE/EN-Launch, 1 Plattform- oder Bildungspilot"),
    ("uof_wb", "Erste Vision-Linie WortBild (Sprachenlernen)", 0.15, ["cl_wortbild"],
     "300 Wörter, Lesehilfe, Sprecher-Audios DE/TR/EN, Didaktik-Prüfung, Kurs-/Klassen-Pilot (M4–M15)"),
    ("uof_legal", "Recht, Marke, Datenschutz", 0.10, ["cl_legal"],
     "Marke LiveFX DE/EU/TR, Beteiligungsvertrag, DSGVO/DSFA Schul-/Familienprofil, Jugendschutz, Steuerberatung"),
    ("uof_res", "Reserve", 0.05, [],
     "Nicht verplant: Verzögerungen bei Einstellung, Store-/Plattform-Freigaben; deckt Infrastruktur-Fixkosten im Fall ohne Umsatz"),
]
uof_first = r
for key, lab, share, lines, what in uof:
    put(M, f"A{r}", lab)
    put(M, f"B{r}", share, F_IN, PCT)
    put(M, f"C{r}", f"=B{r}*{ref('start_cash')}", fmt=EUR, key=f"{key}_amt")
    if lines:
        f = "=" + "+".join(f"'Liquidität 24M'!$AD${rows_cl[l]}" for l in lines)
    else:
        f = "=0"
    put(M, f"D{r}", f, fmt=EUR, key=f"{key}_plan")
    put(M, f"E{r}", f"=C{r}-D{r}", fmt=EUR)
    put(M, f"F{r}", f"=C{r}/18", fmt=EUR)
    put(M, f"G{r}", what, wrap=False)
    reg(f"{key}_share", M, f"B{r}")
    r += 1
put(M, f"A{r}", "Summe", bold=True)
put(M, f"B{r}", f"=SUM(B{uof_first}:B{r-1})", fmt=PCT, bold=True, key="uof_share_total")
put(M, f"C{r}", f"=SUM(C{uof_first}:C{r-1})", fmt=EUR, bold=True, key="uof_total")
put(M, f"D{r}", f"=SUM(D{uof_first}:D{r-1})", fmt=EUR, bold=True, key="uof_plan_total")
put(M, f"E{r}", f"=C{r}-D{r}", fmt=EUR, bold=True)
put(M, f"F{r}", f"=C{r}/18", fmt=EUR, bold=True)
for col in "ABCDEF":
    M[f"{col}{r}"].border = B_TOP
r += 1
put(M, f"A{r}", "Prüfung: Summe Anteile = 100 % und Betrag = Startkasse")
put(M, f"C{r}", f'=IF(AND(ROUND(B{r-1},4)=1,ROUND(C{r-1}-{ref("start_cash")},0)=0),"OK","PRÜFEN")', key="uof_check")
r += 2

section(M, r, "Verschoben in die Folgerunde (nicht aus den 250.000 € finanziert)", 7)
r += 1
for t in ["Story-Engine / Erzählfilm (Linie A): Scene-Director, Figuren, Welten-Packs TR/EN",
          "Studio / Auto-Edit (Linie D): Highlights, Datei-Import, MP4-Export, Pro+",
          "Räume / VR-AR (Linie C): Bühnen-/Klassenzimmer-Preset, WebXR, Brille",
          "Backend/ML-Entwickler:in (Streaming-ASR < 300 ms), Community/Support, 2. Entwickler:in – ab Monat 19 nach Gate 2"]:
    put(M, f"A{r}", "• " + t)
    r += 1
r += 1

section(M, r, "Kostenabgleich Basis: Plan BP §11.3 (500-T€-Fassung) vs. Modell (250-T€-Runde)", 7)
r += 1
header_row(M, r, ["Planzeile BP §11.3", "", "Plan J1", "Modell J1", "Plan J2", "Modell J2", "Was ändert sich"])
r += 1
QS = "'Liquidität 24M'!"
def qsum(rowkeys, col):
    return "+".join(f"{QS}${col}${x}" for x in rowkeys)
cmp_rows = [
    ("Personal: Mobile-/Web-Entwicklung", 75000, 160000, [rows_cl["cl_dev"], rows_cl["cl_tools"], rows_ab["ab_dev"]], None,
     "1 Entwickler:in ab M3 statt sofort; 2. Person erst ab M19"),
    ("Personal: Backend/ML, Community/Support", 0, 60000, [rows_ab["ab_backend"]], None, "erst ab M19 (statt ganz J2)"),
    ("Gründer (Gehalt)", 30000, 48000, [rows_cl["cl_founder"], rows_ab["ab_founder"]], None, "2.000 €/Monat bis M18, danach Plan"),
    ("Freelance Design, Sound, Illustration → WortBild", 15000, 25000, [rows_cl["cl_wortbild"], rows_ab["ab_free"]], None,
     "WortBild 37,5 T€ (M4–M15) statt Freelance-Pauschale"),
    ("Marketing, Creator-Programm, Events, Piloten", 25000, 60000, [rows_cl["cl_mkt1"], rows_cl["cl_mkt2"], rows_cl["cl_pilot"], rows_ab["ab_mkt"]], None,
     "50 T€ in 18 Monaten inkl. 1 Pilot"),
    ("Infrastruktur, KI-API, Store-Gebühren", 8000, 25000, [SR["basis"]["c_inf"]], None, "fix 250 €/Monat + 8 % vom Umsatz"),
    ("Recht, Marke, Steuer, Verwaltung", 12000, 20000, [rows_cl["cl_legal"], rows_ab["ab_legal"]], None, "25 T€ in 18 Monaten (Marke DE/EU/TR, DSFA)"),
]
cmp_first = r
for lab, p1, p2, rows, _, why in cmp_rows:
    put(M, f"A{r}", lab)
    put(M, f"C{r}", p1, F_IN, EUR)
    put(M, f"D{r}", "=" + qsum(rows, "AB"), fmt=EUR)
    put(M, f"E{r}", p2, F_IN, EUR)
    put(M, f"F{r}", "=" + qsum(rows, "AC"), fmt=EUR)
    put(M, f"G{r}", why)
    r += 1
put(M, f"A{r}", "Summe Kosten Basis", bold=True)
for col in "CDEF":
    put(M, f"{col}{r}", f"=SUM({col}{cmp_first}:{col}{r-1})", fmt=EUR, bold=True)
reg("cmp_model_j1", M, f"D{r}")
reg("cmp_model_j2", M, f"F{r}")
r += 1
put(M, f"A{r}", "Kontrolle: = Kosten Basis in der GuV (J1 / J2)")
put(M, f"D{r}", f'=IF(ROUND(D{r-1}-\'GuV 5 Jahre\'!$C${GR["basis"]["cost"]},0)=0,"OK","PRÜFEN")', key="cmp_check1")
put(M, f"F{r}", f'=IF(ROUND(F{r-1}-\'GuV 5 Jahre\'!$D${GR["basis"]["cost"]},0)=0,"OK","PRÜFEN")', key="cmp_check2")
r += 1

# ------------------------------------------------------------------ Blatt: Gates
T = wb.create_sheet("Gates 18M")
widths(T, {"A": 8, "B": 9, "C": 16, "D": 58, "E": 15, "F": 15, "G": 15, "H": 22, "I": 50})
put(T, "A1", "Gates / Meilensteine für 18 Monate – alle Werte sind ZIELE (Basis), Schwelle = konservativ", F_TITLE)
note(T, "A2", "Vorschlag – mit Investor:innen abstimmen. Nutzer/Pro-Ziele aus dem Modell: J1 linear auf den Jahresendwert, J2 linear zwischen den Jahresendwerten.")
r = 4
header_row(T, r, ["Gate", "Monat", "Thema", "KPI", "Ziel Basis", "Schwelle (kons.)", "Optimistisch", "Status", "Gibt frei"])
r += 1

def users_at(s, m):
    u1 = ref(f"users_1_{s}"); u2 = ref(f"users_2_{s}")
    if m <= 12:
        return f"{u1}*{m}/12"
    return f"({u1}+({u2}-{u1})*{m-12}/12)"

def pro_at(s, m):
    return f"{users_at(s, m)}*{ref(f'conv_13_{s}')}"

def mrr_at(s, m):
    return f"INDEX('Liquidität 24M'!$D${SR[s]['rev_rec']}:$AA${SR[s]['rev_rec']},{m})"

gates = [
    ("G0", 3, "Start", "Pro-Abo live (Zahlung), Mobile-/Web-Entwickler:in an Bord", None, "Budget Markteintritt Phase 2"),
    ("G1", 6, "Aktivierung", "Registrierte Nutzer", ("users", 6), "Mobile-App/PWA-Release"),
    ("G1", 6, "Aktivierung", "Zahlende Pro-Nutzer", ("pro", 6), ""),
    ("G1", 6, "Aktivierung", "20 Beta-Creator TR/DE aktiv, 10–20 Botschafter:innen; WortBild mit 300 Wörtern", None, ""),
    ("G1", 6, "Aktivierung", "Woche-4-Retention aktiver Streamer ≥ 30 % (Ziel, erste Kohorten)", None, ""),
    ("G2a", 12, "Traktion", "Registrierte Nutzer", ("users", 12), "DE/EN-Ausbau, Plattform-Gespräche"),
    ("G2a", 12, "Traktion", "Zahlende Pro-Nutzer (Jahresende J1, BP §11.1)", ("pro", 12), ""),
    ("G2a", 12, "Traktion", "Wiederkehrender Umsatz im Monat 12 (MRR, €)", ("mrr", 12), ""),
    ("G2a", 12, "Traktion", "1 Plattform- oder Bildungspilot gestartet; WortBild-Pilot in ≥ 3 Kursen/Klassen", None, ""),
    ("G2", 18, "Folgerunde", "Registrierte Nutzer", ("users", 18), "Seed-Runde, Ausbau-Team ab M19, Story-Engine/Studio"),
    ("G2", 18, "Folgerunde", "Zahlende Pro-Nutzer", ("pro", 18), ""),
    ("G2", 18, "Folgerunde", "Wiederkehrender Umsatz im Monat 18 (MRR, €)", ("mrr", 18), ""),
    ("G2", 18, "Folgerunde", "Monat-3-Retention Pro ≥ 75 %; Conversion Free → Pro ≥ 3 % (Schwelle) / 4 % (Ziel)", None, ""),
    ("G2", 18, "Folgerunde", "WortBild-Pilot ausgewertet; ≥ 5 zahlende Schulen oder Kurse (BP §11.5)", None, ""),
]
gate_rows = []
for gid, m, theme, kpi, metric, frees in gates:
    put(T, f"A{r}", gid, F_BOLD)
    put(T, f"B{r}", m, F_IN, "0")
    put(T, f"C{r}", theme)
    put(T, f"D{r}", kpi)
    if metric:
        kind, mm = metric
        fn = {"users": users_at, "pro": pro_at, "mrr": mrr_at}[kind]
        fmt = EUR if kind == "mrr" else NUM
        put(T, f"E{r}", "=" + fn("basis", mm), fmt=fmt, fill=FILL_KEY, key=f"gate_{kind}_{mm}_basis")
        put(T, f"F{r}", "=" + fn("kons", mm), fmt=fmt, key=f"gate_{kind}_{mm}_kons")
        put(T, f"G{r}", "=" + fn("opt", mm), fmt=fmt, key=f"gate_{kind}_{mm}_opt")
    put(T, f"H{r}", "Ziel (nicht erreicht)")
    put(T, f"I{r}", frees)
    gate_rows.append((gid, m, theme, kpi, metric, frees))
    r += 1
r += 1
note(T, f"A{r}", "Lesart: Gate erfüllt, wenn mindestens die Schwelle (konservatives Szenario) erreicht ist; das Ziel ist das Basis-Szenario.")

# ------------------------------------------------------------------ Blatt: Bewertung
V = wb.create_sheet("Bewertung")
widths(V, {"A": 50, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 70})
put(V, "A1", "Bewertung Pre-Seed – drei Methoden (Vorschlag – mit Steuer-/Rechtsberater prüfen)", F_TITLE)
note(V, "A2", "Alle Eingaben blau; Quellen in Spalte G. Keine Umsätze, Produkt v2.1 funktionsfähig; Solo-Gründer + Investor Relations.")
r = 4
section(V, r, "1. Berkus-Methode (5 Faktoren × max. 500.000 €)", 7)
r += 1
header_row(V, r, ["Faktor", "Max. €", "Score (0–1)", "Wert €", "", "", "Begründung (ehrlich)"])
r += 1
berkus = [
    ("Idee / Grundwert (Problem, Markt)", 0.6, "Neuer Auslöser Stimme → Effekt; großer Markt (BP §6), aber Zahlungsbereitschaft unbewiesen."),
    ("Prototyp / Produkt (Technologierisiko)", 0.7, "v2.1 funktionsfähig: 238 Trigger + 143 freie Sticker, Leistungsmodus (Bildzeit −63 bis −69 %), "
     "sichere GIF-Suche, automatisierte Tests, im Einsatz in eigenen Streams."),
    ("Team (Umsetzungsqualität)", 0.3, "Solo-Gründer + Investor Relations; erste Einstellung offen; Ein-Personen-Risiko (BP §13)."),
    ("Strategische Beziehungen", 0.1, "Noch keine Plattform-, Bildungs- oder Vertriebspartnerschaft unterzeichnet."),
    ("Produkteinführung / Umsatz", 0.1, "Noch kein Umsatz, keine zahlenden Nutzer; Pro-Abo erst ab Monat 3."),
]
b_first = r
for lab, sc, why in berkus:
    put(V, f"A{r}", lab)
    put(V, f"B{r}", 500000, F_IN, EUR)
    put(V, f"C{r}", sc, F_IN, "0.00")
    put(V, f"D{r}", f"=B{r}*C{r}", fmt=EUR)
    put(V, f"G{r}", why)
    r += 1
put(V, f"A{r}", "Berkus-Wert (Pre-Money)", bold=True)
put(V, f"B{r}", f"=SUM(B{b_first}:B{r-1})", fmt=EUR)
put(V, f"D{r}", f"=SUM(D{b_first}:D{r-1})", fmt=EUR, bold=True, fill=FILL_KEY, key="val_berkus")
note(V, f"G{r}", "Methode: Dave Berkus, je Faktor bis 0,5 Mio. € (Pre-Revenue).")
r += 2

section(V, r, "2. Scorecard-Methode (Payne) – Referenz: Median-Pre-Money Pre-Seed DACH/Türkiye", 7)
r += 1
put(V, f"A{r}", "Referenz Median-Pre-Money Pre-Seed (pre-revenue, DACH/TR)")
put(V, f"B{r}", 1500000, F_IN, EUR, FILL_KEY, key="sc_ref")
note(V, f"G{r}", "Abgeleitet: DE Pre-Seed typ. 0,5–1,5 Mio. € (upxcale.de), DE 1–5 Mio. € (capvisory.de); Europa Median Q1 2025 "
     "4,57 Mio. USD (Equidam, institutionelle Runden, daher zu hoch für Solo-pre-revenue); Türkiye niedriger → 1,5 Mio. € (Schätzung).")
r += 1
header_row(V, r, ["Faktor", "Gewicht", "Vergleich (100 % = Median)", "Beitrag", "", "", "Begründung"])
r += 1
scorecard = [
    ("Team", 0.30, 0.60, "Solo-Gründer, Hires offen; Produkt- und Community-Erfahrung als Streamer/Autor"),
    ("Marktgröße / Chance", 0.25, 1.25, "Live-Streaming 97–157 Mrd. USD, SAM 1–3 Mio. Creator (BP §6); Vision-Märkte zusätzlich"),
    ("Produkt / Technologie", 0.15, 1.20, "funktionsfähige v2.1, lokale Verarbeitung, Tests; geringe Grenzkosten"),
    ("Wettbewerbsumfeld", 0.10, 0.90, "kein direkter Sprach-Trigger-Wettbewerber, aber Plattformen können nachbauen"),
    ("Marketing / Vertrieb / Partner", 0.10, 0.70, "eigene Streams + TR-Community; noch keine Partner, kein Vertrieb"),
    ("Bedarf weiterer Finanzierung", 0.05, 1.00, "schlanke Runde mit Gates; Folgerunde nötig (siehe Liquidität)"),
    ("Sonstiges (IP, Bildung, Datenschutz lokal)", 0.05, 1.10, "WortBild/Bildung als zweites Standbein, offline-fähig"),
]
s_first = r
for lab, w, sc, why in scorecard:
    put(V, f"A{r}", lab)
    put(V, f"B{r}", w, F_IN, PCT)
    put(V, f"C{r}", sc, F_IN, PCT)
    put(V, f"D{r}", f"=B{r}*C{r}", fmt="0.000")
    put(V, f"G{r}", why)
    r += 1
put(V, f"A{r}", "Summe Gewichte / Faktor")
put(V, f"B{r}", f"=SUM(B{s_first}:B{r-1})", fmt=PCT, key="sc_wsum")
put(V, f"D{r}", f"=SUM(D{s_first}:D{r-1})", fmt="0.000", key="sc_factor")
r += 1
put(V, f"A{r}", "Scorecard-Wert (Pre-Money)", bold=True)
put(V, f"D{r}", f"={ref('sc_ref')}*D{r-1}", fmt=EUR, bold=True, fill=FILL_KEY, key="val_scorecard")
r += 2

section(V, r, "3. VC-Methode (Exit Jahr 5)", 7)
r += 1
vc = [
    ("vc_rev5", "Umsatz J5 (Basis, aus GuV)", f"='GuV 5 Jahre'!$G${GR['basis']['rev']}", EUR, None, "GuV 5 Jahre, Basis J5"),
    ("vc_mult", "Exit-Multiple (Umsatz)", 4.0, MULT, F_IN,
     "SaaS Capital 2025: private B2B-SaaS 4,8x (bootstrapped) / 5,3x (VC-finanziert); unter 2 Mio. USD ARR 2–3,5x. "
     "Abschlag für B2C/Creator-Tool (höherer Churn) → 4,0x (konservativ)."),
    ("vc_exit", "Exit-Wert Jahr 5", "=B{r0}*B{r1}", EUR, None, "Umsatz × Multiple"),
    ("vc_ret", "Ziel-Rendite (Multiple auf Einsatz)", 15, MULT, F_IN, "Pre-Seed üblich 10–20x; Mittelwert 15x"),
    ("vc_dil", "Verwässerung durch spätere Runden (Seed, Series A, ESOP)", 0.45, PCT, F_IN,
     "Carta 2025: Median Seed ≈ 20 %, Series A ≈ 18 %; + ESOP ≈ 10 % → 1 − 0,8 × 0,82 × 0,9 ≈ 41 %; Annahme 45 %"),
    ("vc_post", "Post-Money heute = Exit ÷ Ziel-Rendite × (1 − Verwässerung)", "=B{r2}/B{r3}*(1-B{r4})", EUR, None, ""),
    ("vc_invest", "Investition (Pre-Seed)", f"={ref('start_cash')}", EUR, None, ""),
    ("val_vc", "VC-Wert (Pre-Money) = Post-Money − Investition", "=MAX(0,B{r5}-B{r6})", EUR, None, "Untergrenze 0"),
]
vc_rows = {}
base = r
for i, (key, lab, val, fmt, font, src) in enumerate(vc):
    vc_rows[i] = base + i
for i, (key, lab, val, fmt, font, src) in enumerate(vc):
    rr = base + i
    if isinstance(val, str):
        val = val.format(**{f"r{k}": vc_rows[k] for k in vc_rows})
    put(V, f"A{rr}", lab, bold=(key == "val_vc"))
    put(V, f"B{rr}", val, font, fmt, FILL_KEY if key == "val_vc" else None, key=key, bold=(key == "val_vc"))
    note(V, f"G{rr}", src)
r = base + len(vc) + 1
put(V, f"A{r}", "VC-Methode: Pre-Money bei Exit-Multiple (Zeilen) × Ziel-Rendite (Spalten)", F_BOLD)
r += 1
rets = [10, 15, 20]
mults = [3, 4, 5, 6]
put(V, f"A{r}", "Multiple ↓ / Rendite →")
for k, rt in enumerate(rets):
    put(V, f"{L(2+k)}{r}", rt, F_IN, MULT)
hdr = r
r += 1
for mlt in mults:
    put(V, f"A{r}", mlt, F_IN, MULT)
    for k in range(len(rets)):
        col = L(2 + k)
        put(V, f"{col}{r}", f"=MAX(0,{ref('vc_rev5')}*$A{r}/{col}${hdr}*(1-{ref('vc_dil')})-{ref('vc_invest')})", fmt=EUR)
    r += 1
grid_first, grid_last = hdr + 1, r - 1
reg("vc_grid_min", V, f"B{grid_first}")
r += 1

section(V, r, "4. Ergebnis und Empfehlung", 7)
r += 1
header_row(V, r, ["Methode", "Pre-Money €", "Gewicht", "", "", "", "Warum dieses Gewicht"])
r += 1
res_first = r
for lab, k, w, why in [("Berkus", "val_berkus", 0.4, "pre-revenue-Standard, Produktstand gut abbildbar"),
                       ("Scorecard", "val_scorecard", 0.4, "Marktvergleich DACH/TR"),
                       ("VC-Methode", "val_vc", 0.2, "hängt stark an 5-Jahres-Prognose ohne Umsatzhistorie → geringeres Gewicht")]:
    put(V, f"A{r}", lab)
    put(V, f"B{r}", f"={ref(k)}", fmt=EUR)
    put(V, f"C{r}", w, F_IN, PCT)
    put(V, f"G{r}", why)
    r += 1
res_last = r - 1
put(V, f"A{r}", "Spanne (Minimum)")
put(V, f"B{r}", f"=MIN(B{res_first}:B{res_last})", fmt=EUR, key="val_min"); r += 1
put(V, f"A{r}", "Spanne (Maximum)")
put(V, f"B{r}", f"=MAX(B{res_first}:B{res_last})", fmt=EUR, key="val_max"); r += 1
put(V, f"A{r}", "Gewichteter Mittelwert")
put(V, f"B{r}", f"=SUMPRODUCT(B{res_first}:B{res_last},C{res_first}:C{res_last})/SUM(C{res_first}:C{res_last})", fmt=EUR, key="val_wavg"); r += 1
put(V, f"A{r}", "Empfohlene Pre-Money-Bewertung (auf 50 T€ gerundet)", bold=True)
put(V, f"B{r}", f"=ROUND({ref('val_wavg')}/50000,0)*50000", fmt=EUR, bold=True, fill=FILL_KEY, key="val_reco"); r += 1
put(V, f"A{r}", "Investition")
put(V, f"B{r}", f"={ref('start_cash')}", fmt=EUR); r += 1
put(V, f"A{r}", "Post-Money = Pre-Money + 250.000 €", bold=True)
put(V, f"B{r}", f"={ref('val_reco')}+{ref('start_cash')}", fmt=EUR, bold=True, key="val_post"); r += 1
put(V, f"A{r}", "Anteil Investor:innen (250.000 € ÷ Post-Money)", bold=True)
put(V, f"B{r}", f"={ref('start_cash')}/{ref('val_post')}", fmt=PCT, bold=True, fill=FILL_KEY, key="val_stake"); r += 1
put(V, f"A{r}", "Anteil Gründer nach der Runde (ohne ESOP)")
put(V, f"B{r}", f"=1-{ref('val_stake')}", fmt=PCT, key="val_founder"); r += 2

section(V, r, "5. Alternative: Wandeldarlehen / SAFE (Vorschlag – mit Steuer-/Rechtsberater prüfen)", 7)
r += 1
put(V, f"A{r}", "Valuation Cap (Pre-Money) ≈ empfohlene Bewertung × Faktor")
put(V, f"C{r}", 1.2, F_IN, '0.00"x"')
put(V, f"B{r}", f"=ROUND({ref('val_reco')}*C{r}/50000,0)*50000", fmt=EUR, fill=FILL_KEY, key="cap")
note(V, f"G{r}", "Cap leicht über dem Preis der Eigenkapital-Variante, weil die Wandlung erst in der Seed-Runde erfolgt; Spanne 1,0–1,2 Mio. €.")
r += 1
put(V, f"A{r}", "Discount auf den Preis der Folgerunde")
put(V, f"B{r}", 0.20, F_IN, PCT, key="discount")
note(V, f"G{r}", "DACH marktüblich 15–25 % (z. B. lexr.com, vektora.eu); Vorschlag 20 %, verhandelbar 15–20 %.")
r += 1
put(V, f"A{r}", "Zins p. a. (wandelt mit)")
put(V, f"B{r}", 0.05, F_IN, PCT, key="cla_interest")
r += 1
put(V, f"A{r}", "Laufzeit bis Wandlung (Monate)")
put(V, f"B{r}", 18, F_IN, "0", key="cla_months")
r += 1
put(V, f"A{r}", "Wandlungsbetrag inkl. Zins")
put(V, f"B{r}", f"={ref('start_cash')}*(1+{ref('cla_interest')}*{ref('cla_months')}/12)", fmt=EUR, key="cla_amount")
r += 1
put(V, f"A{r}", "Beispiel: Pre-Money der Seed-Runde")
put(V, f"B{r}", 3000000, F_IN, EUR, key="seed_pre_example")
r += 1
put(V, f"A{r}", "Wandlungsbewertung = MIN(Cap; Seed-Pre × (1 − Discount))")
put(V, f"B{r}", f"=MIN({ref('cap')},{ref('seed_pre_example')}*(1-{ref('discount')}))", fmt=EUR, key="cla_conv_val")
r += 1
put(V, f"A{r}", "Anteil bei Wandlung (vereinfacht, vor Seed-Geld)", bold=True)
put(V, f"B{r}", f"={ref('cla_amount')}/({ref('cla_conv_val')}+{ref('cla_amount')})", fmt=PCT, bold=True, key="cla_stake")
r += 1
put(V, f"A{r}", "Max. Anteil am Cap ohne Zins (250.000 ÷ (Cap + 250.000))")
put(V, f"B{r}", f"={ref('start_cash')}/({ref('cap')}+{ref('start_cash')})", fmt=PCT, key="cap_stake")
r += 2
note(V, f"A{r}", "Alle Bewertungen: Vorschlag – mit Steuer-/Rechtsberater prüfen. Keine Anlageberatung.")

# ------------------------------------------------------------------ Blatt: Sensitivität
S = wb.create_sheet("Sensitivität")
widths(S, {"A": 34, "B": 11, "C": 11, "D": 11})
for i in range(5, 40):
    S.column_dimensions[L(i)].width = 12
put(S, "A1", "Sensitivität (Basis-Szenario): Umsatz J3 und Runway", F_TITLE)
note(S, "A2", "Conversion ±1 Pkt., ARPU ±20 %, Nutzer ±30 % – je Fall alle Jahre; Kosten wie Basis (Infrastruktur variabel mit Umsatz).")
r = 4
cases = [("Basis", 0, 1, 1), ("Conversion −1 Pkt.", -0.01, 1, 1), ("Conversion +1 Pkt.", 0.01, 1, 1),
         ("ARPU −20 %", 0, 0.8, 1), ("ARPU +20 %", 0, 1.2, 1), ("Nutzer −30 %", 0, 1, 0.7), ("Nutzer +30 %", 0, 1, 1.3)]
hdrs = ["Fall", "Δ Conversion", "ARPU-Faktor", "Nutzer-Faktor", "Umsatz J3 €", "Δ zu Basis", "Wiederk. J1 €", "Wiederk. J2 €",
        "s (Dez J1)", "g (Zuwachs J2)", "Runway (Monate)", "Cash-out-Monat", "Kasse M12 €", "Kasse M18 €", "Kasse M24 €", "Min. Kasse €"]
header_row(S, r, hdrs)
for i in range(5, 17):
    S.column_dimensions[L(i)].width = 14
r += 1
CASE_ROW = {}
for name, dc, af, uf in cases:
    CASE_ROW[name] = r
    r += 1
# Monatsblock
r += 1
MB = r
put(S, f"A{MB}", "Monat", bold=True)
SMC = {m: L(4 + m) for m in range(1, 25)}  # E..AB
for m in range(1, 25):
    put(S, f"{SMC[m]}{MB}", m, F_BOLD, "0")
r += 1
put(S, f"A{r}", "Fixkosten Basis (Pre-Seed + Ausbau + Infrastruktur fix)")
FIX = r
for m in range(1, 25):
    qc = MC[m]
    put(S, f"{SMC[m]}{r}", f"='Liquidität 24M'!{qc}${ROW_PRESEED}+'Liquidität 24M'!{qc}${ROW_AUSBAU}*{ref('ausbau_factor_basis')}+{ref('infra_fix')}", fmt=EUR)
r += 2
CB = {}
for name, dc, af, uf in cases:
    CB[name] = {"rev": r, "cash": r + 1, "neg": r + 2}
    put(S, f"A{r}", f"{name}: Umsatz")
    put(S, f"A{r+1}", f"{name}: Kasse", bold=True)
    put(S, f"A{r+2}", f"{name}: Kasse < 0", F_NOTE)
    r += 4

gb = GR["basis"]
def gcell(col, row):
    return f"'GuV 5 Jahre'!${col}${row}"
for name, dc, af, uf in cases:
    rr = CASE_ROW[name]
    put(S, f"A{rr}", name, F_BOLD)
    put(S, f"B{rr}", dc, F_IN, '+0.0%;-0.0%;0.0%')
    put(S, f"C{rr}", af, F_IN, "0.00")
    put(S, f"D{rr}", uf, F_IN, "0.00")
    def rec_year(col, avgref):
        users = gcell(col, gb["users"])
        conv = gcell(col, gb["conv"])
        return (f"({users}*$D{rr})*({conv}+$B{rr})*{avgref}*{ref('arpu_basis')}*$C{rr}*12"
                f"+({users}*$D{rr})*{ref('pack_share_basis')}*{ref('pack_price')}+{gcell(col, gb['r_ag'])}")
    put(S, f"E{rr}", "=" + rec_year("E", ref("avg_j2")) + f"+{gcell('E', gb['r_b2b'])}", fmt=EUR, key=f"sens_rev3_{name}")
    put(S, f"F{rr}", f"=E{rr}/$E${CASE_ROW['Basis']}-1", fmt='+0.0%;-0.0%;0.0%', key=f"sens_rev3d_{name}")
    put(S, f"G{rr}", "=" + rec_year("C", ref("avg_j1")), fmt=EUR)
    put(S, f"H{rr}", "=" + rec_year("D", ref("avg_j2")), fmt=EUR)
    put(S, f"I{rr}", f"=G{rr}*11.5/72", fmt=EUR)
    put(S, f"J{rr}", f"=(H{rr}-12*I{rr})/78", fmt=EUR)
    cb = CB[name]
    cr = f"$E${cb['cash']}:$AB${cb['cash']}"
    put(S, f"K{rr}", f'=IFERROR(MATCH(1,$E${cb["neg"]}:$AB${cb["neg"]},0)-1,"> 24")', fmt="0", key=f"sens_runway_{name}")
    put(S, f"L{rr}", f'=IFERROR(MATCH(1,$E${cb["neg"]}:$AB${cb["neg"]},0),"keiner")', fmt="0")
    put(S, f"M{rr}", f"=INDEX({cr},12)", fmt=EUR, key=f"sens_cash12_{name}")
    put(S, f"N{rr}", f"=INDEX({cr},18)", fmt=EUR, key=f"sens_cash18_{name}")
    put(S, f"O{rr}", f"=INDEX({cr},24)", fmt=EUR, key=f"sens_cash24_{name}")
    put(S, f"P{rr}", f"=MIN({cr})", fmt=EUR, key=f"sens_cashmin_{name}")
    for m in range(1, 25):
        c = SMC[m]
        mm = f"{c}${MB}"
        put(S, f"{c}{cb['rev']}",
            f"=IF({mm}<=12,$G${rr}*({mm}-0.5)/72,$I${rr}+$J${rr}*({mm}-12))"
            f"+IF({mm}=12+{ref('b2b_month')},{gcell('D', gb['r_b2b'])},0)+IF({mm}={ref('b2b_month')},{gcell('C', gb['r_b2b'])},0)", fmt=EUR)
        prev = ref("start_cash") if m == 1 else f"{SMC[m-1]}{cb['cash']}"
        put(S, f"{c}{cb['cash']}", f"={prev}+{c}{cb['rev']}*(1-{ref('infra_var')})-{c}${FIX}", fmt=EUR, bold=True)
        put(S, f"{c}{cb['neg']}", f"=IF({c}{cb['cash']}<0,1,0)", F_NOTE, "0")
r += 1
put(S, f"A{r}", "Kontrolle: Basis-Fall Kasse M24 = Liquidität Basis M24")
put(S, f"E{r}", f'=IF(ROUND(O{CASE_ROW["Basis"]}-{ref("q_cash_m24_basis")},0)=0,"OK","PRÜFEN")', key="sens_check")
put(S, f"F{r}", f'=IF(ROUND(E{CASE_ROW["Basis"]}-{ref("guv_rev_basis_3")},0)=0,"OK","PRÜFEN")', key="sens_check2")
S.freeze_panes = "B5"

# ------------------------------------------------------------------ Blatt: Diagramme
D = wb.create_sheet("Diagramme")
put(D, "A1", "Diagramme", F_TITLE)
# Datentabelle für Diagramm 1 (Verweise)
put(D, "A3", "Basis (€)", F_BOLD)
put(D, "B3", "Umsatz", F_BOLD)
put(D, "C3", "Kosten", F_BOLD)
put(D, "D3", "EBITDA", F_BOLD)
for j in range(5):
    rr = 4 + j
    put(D, f"A{rr}", YEARS[j])
    put(D, f"B{rr}", f"='GuV 5 Jahre'!{L(3+j)}{gb['rev']}", fmt=EUR)
    put(D, f"C{rr}", f"='GuV 5 Jahre'!{L(3+j)}{gb['cost']}", fmt=EUR)
    put(D, f"D{rr}", f"='GuV 5 Jahre'!{L(3+j)}{gb['ebitda']}", fmt=EUR)
put(D, "F3", "Monat", F_BOLD)
put(D, "G3", "Konservativ", F_BOLD)
put(D, "H3", "Basis", F_BOLD)
put(D, "I3", "Optimistisch", F_BOLD)
put(D, "J3", "Ohne Umsatz", F_BOLD)
for m in range(1, 25):
    rr = 3 + m
    put(D, f"F{rr}", m, fmt="0")
    for k, s in enumerate(SCEN):
        put(D, f"{L(7+k)}{rr}", f"='Liquidität 24M'!{MC[m]}{SR[s]['cash']}", fmt=EUR)
    put(D, f"J{rr}", f"='Liquidität 24M'!{MC[m]}{NR['cash']}", fmt=EUR)
widths(D, {"A": 12, "B": 14, "C": 14, "D": 14, "F": 8, "G": 14, "H": 14, "I": 14, "J": 14})

ch1 = BarChart()
ch1.type = "col"
ch1.title = "Umsatz und Kosten pro Jahr – Basis (€)"
ch1.y_axis.title = "€"
ch1.y_axis.numFmt = '#,##0'
ch1.add_data(Reference(D, min_col=2, max_col=3, min_row=3, max_row=8), titles_from_data=True)
ch1.set_categories(Reference(D, min_col=1, min_row=4, max_row=8))
ch1.height, ch1.width = 9, 18
D.add_chart(ch1, "L3")

ch2 = LineChart()
ch2.title = "Kassenbestand 24 Monate – drei Szenarien und ohne Umsatz (€)"
ch2.y_axis.title = "€"
ch2.x_axis.title = "Monat"
ch2.y_axis.numFmt = '#,##0'
ch2.add_data(Reference(D, min_col=7, max_col=10, min_row=3, max_row=27), titles_from_data=True)
ch2.set_categories(Reference(D, min_col=6, min_row=4, max_row=27))
ch2.height, ch2.width = 10, 18
D.add_chart(ch2, "L23")

# ------------------------------------------------------------------ Blatt: Übersetzung
U = wb.create_sheet("Übersetzung")
widths(U, {"A": 46, "B": 46, "C": 46})
put(U, "A1", "Übersetzung der wichtigsten Begriffe / Çeviri / Translation", F_TITLE)
header_row(U, 3, ["Deutsch", "Türkçe", "English"])
terms = [
    ("Annahmen", "Varsayımlar", "Assumptions"),
    ("GuV 5 Jahre", "5 yıllık gelir tablosu", "5-year P&L"),
    ("Mittelverwendung", "Fonların kullanımı", "Use of funds"),
    ("Liquidität monatlich, 24 Monate", "Aylık nakit akışı, 24 ay", "Monthly cash flow, 24 months"),
    ("Gates / Meilensteine (Ziele)", "Kilometre taşları / hedefler", "Gates / milestones (targets)"),
    ("Bewertung", "Değerleme", "Valuation"),
    ("Sensitivität", "Duyarlılık analizi", "Sensitivity"),
    ("Konservativ / Basis / Optimistisch", "Temkinli / Baz / İyimser", "Conservative / Base / Optimistic"),
    ("Registrierte Nutzer", "Kayıtlı kullanıcılar", "Registered users"),
    ("Conversion Free → Pro", "Ücretsizden Pro'ya dönüşüm", "Free-to-Pro conversion"),
    ("Zahlende Pro-Nutzer", "Ödeme yapan Pro kullanıcılar", "Paying Pro users"),
    ("Umsatz gesamt", "Toplam gelir", "Total revenue"),
    ("Kosten gesamt", "Toplam maliyet", "Total costs"),
    ("EBITDA (Ergebnis)", "FAVÖK (sonuç)", "EBITDA (result)"),
    ("Kumuliertes Ergebnis", "Kümülatif sonuç", "Cumulative result"),
    ("Kasse Monatsende", "Ay sonu nakit", "Cash at month end"),
    ("Runway (Monate)", "Nakit ömrü (ay)", "Runway (months)"),
    ("Cash-out-Monat", "Nakdin bittiği ay", "Cash-out month"),
    ("Nachhaltiger Break-even", "Kalıcı başabaş noktası", "Sustained break-even"),
    ("Folgerunde (Seed) – Mindestbedarf", "Sonraki tur (Seed) – asgari ihtiyaç", "Follow-on round (seed) – minimum need"),
    ("Team/Produkt", "Ekip/Ürün", "Team/product"),
    ("Markteintritt, Creator-Programm, Piloten", "Pazara giriş, içerik üretici programı, pilotlar", "Market entry, creator programme, pilots"),
    ("Vision-Linie WortBild (Sprachenlernen)", "Vizyon hattı WortBild (dil öğrenimi)", "Vision line WortBild (language learning)"),
    ("Recht, Marke, Datenschutz", "Hukuk, marka, veri koruma", "Legal, trademark, data protection"),
    ("Reserve", "Yedek", "Reserve"),
    ("Pre-Money-Bewertung", "Yatırım öncesi değerleme", "Pre-money valuation"),
    ("Post-Money-Bewertung", "Yatırım sonrası değerleme", "Post-money valuation"),
    ("Anteil Investor:innen", "Yatırımcı payı", "Investor stake"),
    ("Wandeldarlehen / SAFE, Valuation Cap, Discount", "Dönüştürülebilir kredi / SAFE, değerleme tavanı, iskonto",
     "Convertible loan / SAFE, valuation cap, discount"),
    ("Vorschlag – mit Steuer-/Rechtsberater prüfen", "Öneri – vergi/hukuk danışmanıyla kontrol edin",
     "Proposal – review with tax/legal adviser"),
    ("Ziel (nicht erreicht)", "Hedef (henüz ulaşılmadı)", "Target (not yet achieved)"),
    ("Schätzung", "Tahmin", "Estimate"),
]
for i, (de, tr, en) in enumerate(terms):
    put(U, f"A{4+i}", de)
    put(U, f"B{4+i}", tr)
    put(U, f"C{4+i}", en)

# Reihenfolge der Blätter
order = ["Annahmen", "GuV 5 Jahre", "Mittelverwendung", "Liquidität 24M", "Gates 18M", "Bewertung", "Sensitivität",
         "Diagramme", "Übersetzung"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = True
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False,
                              color=c.font.color if c.font else None,
                              size=c.font.size if c.font else 11)

wb.save(XLSX)
print("geschrieben:", XLSX)


# =====================================================================================
def recalc(path):
    script = os.environ.get("LIVEFX_RECALC", DEFAULT_RECALC)
    if os.path.exists(script):
        out = subprocess.run([sys.executable, script, path, "120"], capture_output=True, text=True,
                             cwd=os.path.dirname(script))
        txt = out.stdout.strip()
        try:
            res = json.loads(txt[txt.index("{"):])
        except Exception:
            raise SystemExit("recalc fehlgeschlagen: " + out.stdout + out.stderr)
        if "error" in res:
            raise SystemExit("recalc fehlgeschlagen: " + json.dumps(res))
        return res
    # Fallback: LibreOffice-Roundtrip
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, path],
                   check=True, capture_output=True)
    shutil.copy(os.path.join(tmp, os.path.basename(path)), path)
    return {"status": "roundtrip"}


ERRS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def scan_errors(path):
    wbv = load_workbook(path, data_only=True)
    bad = []
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRS):
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    return bad


if "--no-recalc" in sys.argv:
    sys.exit(0)

res = recalc(XLSX)
print("recalc:", json.dumps(res)[:400])
bad = scan_errors(XLSX)
if bad:
    raise SystemExit("Fehlerwerte gefunden: " + ", ".join(bad[:50]))
print("Fehlerprüfung: keine Fehlerwerte")

wv = load_workbook(XLSX, data_only=True)


def v(key):
    sh, cell = R[key]
    val = wv[sh][cell].value
    if isinstance(val, float):
        return round(val, 4) if abs(val) < 10 else round(val, 2)
    return val


def k(x):  # in T€ (eine Nachkommastelle)
    return None if x is None else round(x / 1000, 1)


out = {
    "meta": {
        "title": "LiveFX – Finanzmodell und Bewertung Pre-Seed 250.000 €",
        "generated": date.today().isoformat(),
        "workbook": "live-fx/business/LiveFX_Finanzmodell.xlsx",
        "generator": "live-fx/business/tools/build-finance.py",
        "currency": "EUR",
        "status": "Schätzung / Vorschlag – mit Steuer-/Rechtsberater prüfen; es gibt noch keine Umsätze",
        "timeline": "Closing Jan 2027 = Monat 1; J1 = 2027 (Start Pro-Abo) … J5 = 2031",
        "product": "LiveFX v2.1 (Leistungsmodus, 143 freie Sticker, sichere GIF-Suche KLIPY/GIPHY)",
    },
    "round": {"amount_eur": v("start_cash"), "months_planned": 18, "instrument_options": ["Eigenkapital (Priced Round)", "Wandeldarlehen / SAFE mit Cap + Discount"]},
    "use_of_funds": [
        {"key": key, "label": lab, "share": v(f"{key}_share"), "amount_eur": v(f"{key}_amt"),
         "scheduled_m1_18_eur": v(f"{key}_plan"), "contents": what}
        for key, lab, share, lines, what in uof
    ],
    "use_of_funds_total_eur": v("uof_total"),
    "use_of_funds_check": v("uof_check"),
    "moved_to_follow_on": ["Story-Engine / Erzählfilm (A)", "Studio / Auto-Edit (D)", "Räume / VR-AR (C)",
                           "Backend/ML, Community/Support, 2. Entwickler:in ab Monat 19"],
    "cost_lines_preseed": [
        {"key": key, "label": label, "eur_per_month": rate, "from_month": frm, "to_month": to,
         "one_off_eur": once, "one_off_month": once_m, "sum_m1_18_eur": v(f"{key}_sum18")}
        for key, label, cat, rate, frm, to, once, once_m, why in cost_lines
    ],
    "runway": {
        "no_revenue_months": v("q_runway_nr"),
        "no_revenue_cash_out_month": v("q_cashout_m_nr"),
        "no_revenue_avg_burn_m1_18_eur": v("q_burn18"),
        "by_scenario": {s: {"runway_months": v(f"q_runway_{s}"), "cash_out_month": v(f"q_cashout_m_{s}"),
                             "sustained_break_even_month": v(f"q_be_{s}"),
                             "min_cash_eur": v(f"q_cash_min_{s}"), "min_cash_month": v(f"q_cash_min_m_{s}"),
                             "below_buffer_month": v(f"q_bufm_{s}")} for s in SCEN},
    },
    "cash": {s: {f"m{m}": v(f"q_cash_m{m}_{s}") for m in (6, 12, 18, 24)} for s in SCEN + ["nr"]},
    "follow_on": {
        "by_scenario": {s: {"buffer_eur": v(f"q_buffer_{s}"), "cash_end_j3_without_round_eur": v(f"q_cash_j3_{s}"),
                             "liquidity_gap_eur": v(f"q_gap_{s}"), "deferred_lines_eur": v(f"q_defer_{s}"),
                             "minimum_need_eur": v(f"q_followon_{s}"), "close_by_month": v(f"q_followon_m_{s}")}
                        for s in SCEN},
        "seed_planning_guide_eur": v("q_seed_guide"),
        "timing": "Seed nach Gate 2 (Monat 15–18), Gespräche ab Monat 12",
    },
    "pnl_5y": {
        s: [{"year": YEARS[j], "revenue_eur": v(f"guv_rev_{s}_{j+1}"), "costs_eur": v(f"guv_cost_{s}_{j+1}"),
             "ebitda_eur": v(f"guv_ebitda_{s}_{j+1}"), "cumulative_eur": v(f"guv_cum_{s}_{j+1}"),
             "cash_year_end_eur": v(f"guv_cash_{s}_{j+1}")} for j in range(5)]
        for s in SCEN
    },
    "wortbild_revenue_memo_eur": [v(f"wortbild_rev_{j}") for j in (1, 2, 3)],
    "valuation": {
        "berkus_eur": v("val_berkus"),
        "berkus_scores": [{"factor": lab, "score": sc, "value_eur": round(500000 * sc)} for lab, sc, _ in berkus],
        "scorecard_eur": v("val_scorecard"),
        "scorecard_reference_eur": v("sc_ref"),
        "scorecard_factor": v("sc_factor"),
        "vc_method": {"revenue_y5_basis_eur": v("vc_rev5"), "exit_multiple": v("vc_mult"), "exit_value_eur": v("vc_exit"),
                       "target_return": v("vc_ret"), "later_dilution": v("vc_dil"), "post_money_today_eur": v("vc_post"),
                       "pre_money_eur": v("val_vc")},
        "range_min_eur": v("val_min"), "range_max_eur": v("val_max"), "weighted_avg_eur": v("val_wavg"),
        "weights": {"berkus": 0.4, "scorecard": 0.4, "vc": 0.2},
        "recommended_pre_money_eur": v("val_reco"),
        "post_money_eur": v("val_post"),
        "investor_stake": v("val_stake"),
        "founder_stake_after": v("val_founder"),
        "convertible": {"valuation_cap_eur": v("cap"), "discount": v("discount"), "interest_pa": v("cla_interest"),
                         "months_to_conversion": v("cla_months"), "conversion_amount_eur": v("cla_amount"),
                         "example_seed_pre_money_eur": v("seed_pre_example"), "conversion_valuation_eur": v("cla_conv_val"),
                         "stake_at_conversion": v("cla_stake"), "max_stake_at_cap_no_interest": v("cap_stake"),
                         "discount_range_note": "15–20 % verhandelbar"},
        "note": "Vorschlag – mit Steuer-/Rechtsberater prüfen",
    },
    "gates": [],
    "sensitivity_basis": [
        {"case": name, "revenue_j3_eur": v(f"sens_rev3_{name}"), "delta_vs_base": v(f"sens_rev3d_{name}"),
         "runway_months": v(f"sens_runway_{name}"), "cash_m12_eur": v(f"sens_cash12_{name}"),
         "cash_m18_eur": v(f"sens_cash18_{name}"), "cash_m24_eur": v(f"sens_cash24_{name}"),
         "min_cash_eur": v(f"sens_cashmin_{name}")}
        for name, *_ in cases
    ],
    "checks": {"use_of_funds": v("uof_check"), "cost_reconciliation": [v("cmp_check1"), v("cmp_check2")],
               "sensitivity_base_equals_model": [v("sens_check"), v("sens_check2")]},
    "assumption_changes": changes,
    "sources": [
        {"id": "BP", "text": "BUSINESSPLAN.md §6, §9, §11.1–11.6, §12–14 (LiveFX, Okt. 2026) – Nutzer, Conversion, ARPU, Packs, Agentur, B2B, Kosten"},
        {"id": "MEMO", "text": "investor/INVESTOR-MEMO.de.md §10–11 (500-T€-Fassung) – Aufteilung Vision-Budget, Gates"},
        {"id": "CHANGELOG", "text": "live-fx/CHANGELOG.md 2.1.0 – Produktstand v2.1"},
        {"id": "QUELLEN", "text": "QUELLEN.md [Quelle 1–60] – Marktzahlen (indirekt über BP)"},
        {"id": "Equidam", "text": "Equidam – Pre-Seed Valuations Q1 2025: US & Europe (Europa-Median Pre-Seed 4,57 Mio. USD)",
         "url": "https://www.equidam.com/startup-valuation-delta-q1-2025/"},
        {"id": "upxcale", "text": "upxcale – Pre-Seed Funding 2026: Bewertung, Quellen, Zeitplan (Pre-Seed DE typ. 0,5–1,5 Mio. € Pre-Money; DACH 1,5–5 Mio. €)",
         "url": "https://upxcale.de/blog/pre-seed-funding/"},
        {"id": "capvisory", "text": "Capvisory – Startup Funding Stages with Benchmarks 2025 (Pre-Seed DE 1–5 Mio. €)",
         "url": "https://capvisory.de/the-startup-funding-stages-from-pre-seed-to-series-c/"},
        {"id": "SaaSCapital", "text": "SaaS Capital – 2025 Private SaaS Company Valuations (4,8x bootstrapped / 5,3x equity-backed)",
         "url": "https://www.saas-capital.com/blog-posts/private-saas-company-valuations-multiples/"},
        {"id": "Carta", "text": "Carta – State of Private Markets Q1 2025 (Median-Verwässerung Series A 17,9 %; Seed ≈ 20 %)",
         "url": "https://carta.com/data/state-of-private-markets-q1-2025/"},
        {"id": "lexr", "text": "Lexr – Das Wandeldarlehen in der Praxis (Discount, Cap, Zins)",
         "url": "https://www.lexr.com/en-de/blog/convertible-loan-in-practice-conversion-interest-rate-discount-cap-valuation/"},
        {"id": "vektora", "text": "Vektora – Pre-Seed-Finanzierung in Deutschland (Wandeldarlehen mit 15–25 % Discount üblich)",
         "url": "https://vektora.eu/de/fachbeitraege/pre-seed-finanzierung-in-deutschland-instrumente-und-prozess"},
        {"id": "note", "text": "Web-Recherche 04.10.2026 über Suchergebnisse; die Seiten selbst waren aus der Arbeitsumgebung nicht abrufbar – Werte vor Verwendung gegenprüfen."},
    ],
}
for gid, m, theme, kpi, metric, frees in gate_rows:
    item = {"gate": gid, "month": m, "theme": theme, "kpi": kpi, "status": "Ziel (nicht erreicht)", "releases": frees}
    if metric:
        kind, mm = metric
        item.update({"target_basis": v(f"gate_{kind}_{mm}_basis"), "threshold_kons": v(f"gate_{kind}_{mm}_kons"),
                     "optimistic": v(f"gate_{kind}_{mm}_opt")})
    out["gates"].append(item)

with open(JSON_OUT, "w", encoding="utf-8") as fh:
    json.dump(out, fh, ensure_ascii=False, indent=2)
print("geschrieben:", JSON_OUT)
